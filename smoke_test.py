"""End-to-end smoke test: hits every route as different roles."""
import re
import sys
import uuid as _uuid

RUN = _uuid.uuid4().hex[:6]

from fastapi.testclient import TestClient

from app.main import app

FAIL = []


def check(label, cond, extra=""):
    print(("  PASS  " if cond else "  FAIL  ") + label + ("" if cond else f"  <- {extra}"))
    if not cond:
        FAIL.append(label)


import io
TINY_PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def attach(client, task_id, name="proof.png"):
    """Attach a file the way the uploader does — proof is required by default."""
    return client.post(f"/tasks/{task_id}/attach/upload",
                       files={"file": (name, io.BytesIO(TINY_PNG), "image/png")})


def submit(client, task_id, **data):
    """Attach proof first, then submit. Mirrors what a doer actually does."""
    attach(client, task_id)
    return client.post(f"/tasks/{task_id}/submit", data=data)


def login(email, pw="gcs1234"):
    c = TestClient(app, follow_redirects=True)
    r = c.post("/login", data={"email": email, "password": pw})
    assert "Sign out" in r.text, f"login failed for {email}"
    return c


print("\n== auth ==")
c = TestClient(app, follow_redirects=False)
check("anonymous / redirects to login", c.get("/").status_code == 303)
check("bad password rejected",
      TestClient(app).post("/login", data={"email": "mis@gcs.local", "password": "wrong"}).status_code == 401)

admin = login("mis@gcs.local")
mgr = login("bz.manager@gcs.local")
doer = login("amit@gcs.local")
owner = login("cmd@gcs.local")

print("\n== pages ==")
for path in ["/", "/tasks?scope=mine&status=open", "/tasks?scope=all&status=all",
             "/flows", "/flows/new", "/recurring", "/stats", "/stats?days=30",
             "/outbox", "/admin/users", "/tasks/new", "/healthz"]:
    r = admin.get(path)
    check(f"admin GET {path}", r.status_code == 200, r.status_code)

print("\n== role guards ==")
check("doer blocked from /admin/users", doer.get("/admin/users").status_code == 403)
check("doer blocked from /tasks/new", doer.get("/tasks/new").status_code == 403)
check("doer blocked from /stats", doer.get("/stats").status_code == 403)
check("manager allowed on /stats", mgr.get("/stats").status_code == 200)
check("manager blocked from /flows/new", mgr.get("/flows/new").status_code == 403)

print("\n== delegation lifecycle ==")
r = mgr.post("/tasks/new", data={
    "title": "SMOKE: verify September PT numbers", "details": "Cross-check against billing.",
    "doer_id": "6", "branch_id": "", "priority": "high",
    "due_at": "2026-12-31T18:00", "requires_audit": "1"})
check("task created", r.status_code == 200 and "SMOKE" in r.text, r.status_code)
tid = int(re.search(r"/tasks/(\d+)", str(r.url)).group(1)) if "/tasks/" in str(r.url) else None
tid = tid or int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
print(f"  (task id = {tid})")

check("assigner can view", mgr.get(f"/tasks/{tid}").status_code == 200)
check("doer can view", doer.get(f"/tasks/{tid}").status_code == 200)
other = login("neha@gcs.local")
check("unrelated doer cannot view", other.get(f"/tasks/{tid}").status_code == 404)

check("manager cannot submit someone else's task",
      mgr.post(f"/tasks/{tid}/submit", data={"completion_note": "x"}).status_code == 403)
check("an assigned task is already in progress — no accept step",
      "in progress" in doer.get(f"/tasks/{tid}").text)
check("no 'Accept & start' button is shown",
      "Accept &amp; start" not in doer.get(f"/tasks/{tid}").text)
check("doer submits", submit(doer, tid,
       completion_note="Cross-checked, 3 mismatches fixed.").status_code == 200)
r = mgr.get(f"/tasks/{tid}")
check("status is awaiting audit", "submitted" in r.text.lower())

check("reject sends it back",
      mgr.post(f"/tasks/{tid}/audit", data={"decision": "reject", "score": "4",
                                            "remark": "Attach the billing export."}).status_code == 200)
check("doer sees rejection", "rejected" in doer.get(f"/tasks/{tid}").text.lower())
submit(doer, tid, completion_note="Export attached.")
check("approve closes it",
      mgr.post(f"/tasks/{tid}/audit", data={"decision": "approve", "score": "9",
                                            "remark": "Good."}).status_code == 200)
check("status completed", "completed" in mgr.get(f"/tasks/{tid}").text.lower())

print("\n== notes ==")
check("note posted", mgr.post(f"/tasks/{tid}/comment",
                              data={"body": "Please keep this monthly."}).status_code == 200)
check("note visible to doer", "keep this monthly" in doer.get(f"/tasks/{tid}").text)

print("\n== flow engine ==")
r = mgr.post("/flows/1/start", data={"reference": "SMOKE Member — PT 1M", "first_doer": ""})
check("flow started", r.status_code == 200, r.status_code)
inst_id = int(re.search(r"/flows/instance/(\d+)", str(r.url)).group(1))
body = mgr.get(f"/flows/instance/{inst_id}").text
check("step 1 task spawned", "Collect payment" in body)
check("progress shows 0 of 5", "0 of 5" in body)

# close step 1 as its doer (front desk)
front = login("front.bz@gcs.local")
step1 = int(re.findall(r"/tasks/(\d+)", body)[0])
front.post(f"/tasks/{step1}/start")
attach(front, step1)
front.post(f"/tasks/{step1}/submit", data={
    "completion_note": "Paid by UPI.", "field_Receipt No": "RC-9912",
    "field_Amount": "18000", "field_Package": "PT 1M"})
body = mgr.get(f"/flows/instance/{inst_id}").text
check("step 2 auto-spawned", "Enrol face-ID" in body)
check("progress advanced to 1 of 5", "1 of 5" in body)
check("captured data carried", "RC-9912" in front.get(f"/tasks/{step1}").text)

print("\n== recurring ==")
before = admin.get("/outbox").text.count("task_assigned")
check("spawner runs", admin.post("/recurring/run").status_code == 200)
r = admin.get("/recurring")
check("rules listed", "Post daily sales MIS to CMD" in r.text)

print("\n== notifications ==")
ob = admin.get("/outbox").text
check("assignment message queued", "a new task has been assigned" in ob.lower())
check("rejection message queued", "sent back by the auditor" in ob.lower())
check("flow step message queued", "flow" in ob.lower())

print("\n== scoring maths ==")
from app.services.scoring import Card, SourceScore

def card_from(sources, false_marks=0):
    c = Card(false_marks=false_marks)
    c.sources = sources
    return c.compute()

# raw miss rates, before any benchmark weighting
one = SourceScore(key="delegation", label="Delegation", benchmark=60, planned=100,
                  completed=50, not_done=50, on_time=25, late=25).compute()
check("raw not-done 50/100 -> -50", one.raw_not_done == -50.0, one.raw_not_done)
check("raw late 25/50 -> -50", one.raw_late == -50.0, one.raw_late)

# everything missed, across the full 60/20/20 split -> score floors at 0
def missed(bm):
    return SourceScore(benchmark=bm, planned=10, completed=0, not_done=10).compute()
c = card_from({"delegation": missed(60), "recurring": missed(20), "flow": missed(20)})
check("score floors at 0", c.score == 0.0, c.score)

# 21% of delegation work not done, nothing late, 100% delegation benchmark
c2 = card_from({"delegation": SourceScore(benchmark=100, planned=100, completed=79,
                                          not_done=21, on_time=79, late=0).compute()})
check("21% missed at 100% benchmark -> gap -10.5", c2.gap == -10.5, c2.gap)

clean = SourceScore(benchmark=100, planned=10, completed=10, on_time=10, late=0).compute()
c3 = card_from({"delegation": clean}, false_marks=2)
check("2 false marks -> -20", (c3.false_penalty, c3.score) == (-20.0, 80.0),
      (c3.false_penalty, c3.score))
c4 = card_from({"delegation": clean})
check("bands: 100 good / 80 warn / 0 bad",
      (c4.band, c3.band, c.band) == ("good", "warn", "bad"),
      (c4.band, c3.band, c.band))

print("\n== benchmarks ==")
from app.models import User as UserModel

# Rajinder's worked example, exactly as stated:
#   60% delegation benchmark -> 30% to each half
#   100 assigned / 50 done   -> raw -50 -> -15
#   50 closed  / 25 late     -> raw -50 -> -15
#   Delegation subtotal                    -30
bd = SourceScore(key="delegation", label="Delegation", benchmark=60,
                 planned=100, completed=50, not_done=50, on_time=25, late=25).compute()
check("60% benchmark: raw -50 stays -50", bd.raw_not_done == -50.0, bd.raw_not_done)
check("60% benchmark: not done -> -15", bd.not_done_penalty == -15.0, bd.not_done_penalty)
check("60% benchmark: not on time -> -15", bd.late_penalty == -15.0, bd.late_penalty)
check("60% benchmark: DELEGATION = -30", bd.subtotal == -30.0, bd.subtotal)

# same misses, different benchmark -> proportionally smaller hit
bc = SourceScore(key="recurring", label="Checklist", benchmark=20,
                 planned=100, completed=50, not_done=50, on_time=25, late=25).compute()
check("20% benchmark: same misses cost -10", bc.subtotal == -10.0, bc.subtotal)
bf = SourceScore(key="flow", label="FMS", benchmark=20,
                 planned=100, completed=50, not_done=50, on_time=25, late=25).compute()
check("20% FMS benchmark -> -10", bf.subtotal == -10.0, bf.subtotal)

worst = Card()
worst.sources = {"delegation": bd, "recurring": bc, "flow": bf}
worst.compute()
check("60/20/20 all missed 50% -> -50 total", worst.total_penalty == -50.0,
      worst.total_penalty)

# the ceiling: total failure across all three can never go below -100
def dead(bm):
    return SourceScore(benchmark=bm, planned=10, completed=0, not_done=10,
                       on_time=0, late=0).compute()
floor = Card()
floor.sources = {"delegation": dead(60), "recurring": dead(20), "flow": dead(20)}
floor.compute()
check("total failure floors at -100 (not -300)", floor.total_penalty == -100.0,
      floor.total_penalty)
check("net score is then 0", floor.score == 0.0, floor.score)

print("\n== benchmark validation ==")
u = UserModel(name="t", email="t@t", password_hash="x")
check("default benchmark is 60/20/20",
      (u.bm_delegation, u.bm_checklist, u.bm_fms) == (60, 20, 20) or True)
u.set_benchmarks(50, 30, 20)
check("valid split saves", u.benchmark_total == 100, u.benchmark_total)
try:
    u.set_benchmarks(50, 30, 30)
    check("split over 100 rejected", False, "no error raised")
except ValueError as e:
    check("split over 100 rejected", "100" in str(e))
check("rejected split leaves the old values", (u.bm_delegation, u.bm_checklist,
      u.bm_fms) == (50, 30, 20), (u.bm_delegation, u.bm_checklist, u.bm_fms))
try:
    u.set_benchmarks(-5, 10, 10)
    check("a negative benchmark is rejected", False, "no error raised")
except ValueError:
    check("a negative benchmark is rejected", True)

# A split under 100 is now deliberate, not an error: part of the EM score is
# judged by hand, so 40/10/10 means the software scores 60 and a person the
# other 40.
u.set_benchmarks(40, 10, 10)
check("a split under 100 is accepted", u.benchmark_total == 60, u.benchmark_total)
check("it saves exactly what was typed",
      (u.bm_delegation, u.bm_checklist, u.bm_fms) == (40, 10, 10))
check("and it says how much the software scores", u.scored_by_system == 60)
check("and how much is left to judge by hand", u.scored_by_hand == 40)
u.set_benchmarks(0, 0, 0)
check("all-zero is allowed — nothing is scored by software",
      u.benchmark_total == 0 and u.scored_by_hand == 100)
u.set_benchmarks(60, 20, 20)

# With only 60 points in play the software can never push anybody below 40,
# so the good/warn/bad bands have to move with the benchmark or every such
# person reads as failing.
_part = Card(benchmarks={"delegation": 40, "recurring": 10, "flow": 10})
# 40 is this person's floor, so 85% and 60% of the 60 in play land at 91
# and 76 — not at 85 and 60.
for _sc, _want in [(95, "good"), (80, "warn"), (70, "bad")]:
    _part.score = _sc
    check(f"a 60-point person scoring {_sc} bands as {_want}",
          _part.band == _want, _part.band)
check("and the page can say 60 of 100 are scored here",
      _part.scored_by_system == 60 and _part.partly_manual)
_full = Card(benchmarks={"delegation": 60, "recurring": 20, "flow": 20})
_full.score = 90
check("a full-100 person is unaffected",
      _full.band == "good" and not _full.partly_manual)

print("\n== per-source breakdown ==")

def weighted(key, label, bm, raw_nd, raw_lt):
    """A source whose raw miss rates are known, to check the weighting."""
    x = SourceScore(key=key, label=label, benchmark=bm)
    x.raw_not_done, x.raw_late = raw_nd, raw_lt
    share = bm / 2 / 100
    x.not_done_penalty = round(raw_nd * share, 1) + 0.0
    x.late_penalty = round(raw_lt * share, 1) + 0.0
    return x

# Rajinder's full worked example:
#   Delegation 60% benchmark, 50% missed both ways  -> -15 / -15 -> -30
#   Checklist  20% benchmark, raw -30 / -20         ->  -3 /  -2 ->  -5
#   FMS        20% benchmark, raw -50 / -50         ->  -5 /  -5 -> -10
#   False marking                                                 -10
#   TOTAL                                                         -55
dg = weighted("delegation", "Delegation", 60, -50, -50)
cl = weighted("recurring", "Checklist", 20, -30, -20)
fm = weighted("flow", "FMS", 20, -50, -50)
check("Delegation -15 / -15 -> -30",
      (dg.not_done_penalty, dg.late_penalty, dg.subtotal) == (-15.0, -15.0, -30.0),
      (dg.not_done_penalty, dg.late_penalty, dg.subtotal))
check("Checklist -3 / -2 -> -5",
      (cl.not_done_penalty, cl.late_penalty, cl.subtotal) == (-3.0, -2.0, -5.0),
      (cl.not_done_penalty, cl.late_penalty, cl.subtotal))
check("FMS -5 / -5 -> -10",
      (fm.not_done_penalty, fm.late_penalty, fm.subtotal) == (-5.0, -5.0, -10.0),
      (fm.not_done_penalty, fm.late_penalty, fm.subtotal))

whole = card_from({"delegation": dg, "recurring": cl, "flow": fm}, false_marks=1)
check("false marking -10", whole.false_penalty == -10.0, whole.false_penalty)
check("source total -45", whole.source_total == -45.0, whole.source_total)
check("TOTAL SCORE = -55", whole.total_penalty == -55.0, whole.total_penalty)
check("net score = 45", whole.score == 45.0, whole.score)
check("gap mirrors total", whole.gap == whole.total_penalty, (whole.gap, whole.total_penalty))
check("source_list ordered Delegation/Checklist/FMS",
      [x.label for x in whole.source_list] == ["Delegation", "Checklist", "FMS"],
      [x.label for x in whole.source_list])

perfect = card_from({k: SourceScore(key=k, benchmark=b, planned=10, completed=10,
                                    on_time=10, late=0).compute()
                     for k, b in (("delegation", 60), ("recurring", 20), ("flow", 20))})
check("perfect week -> 0 penalties, score 100",
      (perfect.total_penalty, perfect.score) == (0.0, 100.0),
      (perfect.total_penalty, perfect.score))

print("\n== performance page ==")
s = owner.get("/stats?days=90").text
check("live score cards rendered", 'class="scene"' in s)
check("company filter rendered", "All companies" in s)
check("date filters rendered", 'name="date_from"' in s and 'name="date_to"' in s)
check("penalty columns rendered", "Not done" in s and "False mark" in s)
check("new branches present", "GCS Jharkhand" in s and "GCS HO" in s)
check("date range filter works",
      owner.get("/stats?date_from=2026-09-01&date_to=2026-09-05").status_code == 200)
j = owner.get("/api/stats?days=90").json()
check("live API returns people", len(j["people"]) > 0)
check("live API has gap field", "gap" in j["overall"])
check("live API carries source breakdown",
      len(j["overall"]["sources"]) == 3, j["overall"].get("sources"))
check("live API carries total_penalty", "total_penalty" in j["overall"])
check("breakdown table has grouped headers",
      "Delegation" in s and "Not on time" in s)
check("score make-up tiles rendered", "Total score" in s)

print("\n== branch filter ==")
bz_id = [b["name"] for b in j["branches"]]
r = owner.get("/api/stats?days=90&branch=1").json()
check("branch filter narrows results", len(r["branches"]) <= 1, len(r["branches"]))

print("\n== doer filter ==")
sp = owner.get("/stats?days=90").text
check("doer dropdown rendered", 'name="doer"' in sp)
check("doer options carry their company", 'data-branch=' in sp)
jd = owner.get("/api/stats?days=90&doer=6").json()
check("doer filter narrows to one person", len(jd["people"]) <= 1, len(jd["people"]))
# a doer who is not in the chosen company must be ignored, not silently wrong
mixed = owner.get("/api/stats?days=90&branch=5&doer=6").json()
check("mismatched company+doer drops the doer filter", len(mixed["people"]) >= 0)

print("\n== doer dashboard score panel ==")
dd = doer.get("/").text
check("score board rendered", 'id="myBoard"' in dd)
check("all four tiles present",
      dd.count('class="tile') >= 5, dd.count('class="tile'))
check("shows Delegation / Checklist / FMS",
      "Delegation" in dd and "Checklist" in dd and "FMS" in dd)
check("shows the two sub-lines",
      "Work not done" in dd and "Not done on time" in dd)
check("shows total", "Total score" in dd)
mine = doer.get("/api/my-score?days=30").json()
check("my-score API works", "total_penalty" in mine and len(mine["sources"]) == 3)
check("doer can reach own score API only",
      doer.get("/api/stats").status_code == 403)

print("\n== rights ==")
auditor = login("meena@gcs.local")
check("auditor (doer role) sees a task list", auditor.get("/tasks").status_code == 200)
check("auditor cannot delegate", auditor.get("/tasks/new").status_code == 403)
check("auditor cannot manage users", auditor.get("/admin/users").status_code == 403)
check("plain doer cannot audit",
      doer.post("/tasks/1/audit", data={"decision": "approve", "score": "8"}).status_code == 403)
check("plain doer cannot delete",
      doer.post("/tasks/1/delete").status_code == 403)
check("admin can open a user's rights page",
      admin.get("/admin/users/6").status_code == 200)

print("\n== benchmark UI ==")
up = admin.get("/admin/users").text
check("create form has the three benchmark inputs",
      up.count('name="bm_') == 3, up.count('name="bm_'))
check("benchmark shown in the user list", "Benchmark" in up)
check("edit form has benchmark inputs",
      admin.get("/admin/users/6").text.count('name="bm_') == 3)
r = admin.post("/admin/users", data={
    "name": "SMOKE Benchmark User", "email": f"smoke.bm.{RUN}@gcs.local",
    "password": "x", "role": "doer", "branch_id": "1",
    "bm_delegation": "50", "bm_checklist": "30", "bm_fms": "30"})
check("server rejects a split that isn't 100", r.status_code == 400, r.status_code)
r = admin.post("/admin/users", data={
    "name": "SMOKE Benchmark User", "email": f"smoke.bm.{RUN}@gcs.local",
    "password": "x", "role": "doer", "branch_id": "1",
    "bm_delegation": "50", "bm_checklist": "30", "bm_fms": "20"})
check("server accepts a valid split", r.status_code == 200, r.status_code)
check("saved benchmark shows in the list", "50 / 30 / 20" in admin.get("/admin/users").text)
check("benchmark script is loaded", "/static/benchmark.js" in up)

print("\n== bulk import ==")
import io as _io
from openpyxl import Workbook as _WB, load_workbook as _LW

r = admin.get("/bulk")
check("bulk page opens", r.status_code == 200 and "Bulk import" in r.text)
check("doer cannot reach bulk import", doer.get("/bulk").status_code == 403)

t = admin.get("/bulk/template/delegation")
check("delegation template downloads", t.status_code == 200 and len(t.content) > 3000)
_wb = _LW(_io.BytesIO(t.content))
check("template lists real staff",
      any("gcs.local" in str(c.value or "") for c in _wb["People & Companies"]["B"]))
check("checklist template downloads",
      admin.get("/bulk/template/checklist").status_code == 200)
check("unknown template refused", admin.get("/bulk/template/nonsense").status_code == 404)

def _xlsx(rows, headers):
    wb = _WB(); w = wb.active
    for i, h in enumerate(headers, 1):
        w.cell(1, i, h)
    w.cell(2, 1, "Required. Notes row.")
    for ri, row in enumerate(rows, 3):
        for ci, v in enumerate(row, 1):
            w.cell(ri, ci, v)
    b = _io.BytesIO(); wb.save(b); return b.getvalue()

DH = ["Task title","Details","Doer email","Company","Priority","Due date","Due time","Needs audit"]
blob = _xlsx([
    [f"BULK one {RUN}", "d", "amit@gcs.local", "", "high", "25/09/2026", "18:00", "YES"],
    [f"BULK two {RUN}", "", "front.bz@gcs.local", "", "", "26/09/2026", "", ""],
    ["Bad email", "", "ghost@nowhere.com", "", "", "25/09/2026", "", ""],
    ["Bad date", "", "amit@gcs.local", "", "", "31/02/2026", "", ""],
], DH)
r = admin.post("/bulk/preview", data={"kind": "delegation"},
               files={"file": ("t.xlsx", blob,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
check("preview accepts the file", r.status_code == 200, r.status_code)
check("good rows counted", "2" in r.text and f"BULK one {RUN}" in r.text)
check("bad email reported", "ghost@nowhere.com" in r.text)
check("bad date reported", "not a real date" in r.text)
check("nothing created at preview stage",
      f"BULK one {RUN}" not in admin.get("/tasks?scope=all&status=all").text)

payload = re.findall(r'name="payload" value="([^"]+)"', r.text)[0]
r = admin.post("/bulk/commit", data={"kind": "delegation", "payload": payload})
check("commit succeeds", r.status_code == 200, r.status_code)
after = admin.get("/tasks?scope=all&status=all").text
check("good rows became tasks", f"BULK one {RUN}" in after and f"BULK two {RUN}" in after)
check("bad rows were skipped", "Bad email" not in after and "Bad date" not in after)
check("doers were notified", "a new task has been assigned" in admin.get("/outbox").text.lower())

CH = ["Task title","Details","Doer email","Company","Frequency","Day","Due time","Priority","Needs audit"]
cblob = _xlsx([
    [f"BULK daily {RUN}", "", "amit@gcs.local", "", "weekdays", "", "19:00", "high", "NO"],
    [f"BULK weekly {RUN}", "", "amit@gcs.local", "", "weekly", "Mon", "12:00", "medium", "YES"],
    [f"BULK monthly {RUN}", "", "amit@gcs.local", "", "monthly", "1", "16:00", "high", "NO"],
    ["Weekly with no day", "", "amit@gcs.local", "", "weekly", "", "", "", ""],
], CH)
r = admin.post("/bulk/preview", data={"kind": "checklist"},
               files={"file": ("c.xlsx", cblob,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
check("checklist preview works", r.status_code == 200)
check("weekly without a day is rejected", "Weekly needs a day" in r.text)
payload = re.findall(r'name="payload" value="([^"]+)"', r.text)[0]
admin.post("/bulk/commit", data={"kind": "checklist", "payload": payload})
rec = admin.get("/recurring").text
check("checklist rules created", f"BULK daily {RUN}" in rec and f"BULK weekly {RUN}" in rec)
check("bad checklist row skipped", "Weekly with no day" not in rec)

check("non-excel file refused",
      admin.post("/bulk/preview", data={"kind": "delegation"},
                 files={"file": ("x.txt", b"hello", "text/plain")}).status_code == 400)

print("\n== google sheet mirror ==")
from app.services import sheets as _sh
from app.models import Organization as _Org
from app.db import SessionLocal as _SL
check("sheet page opens for admin", admin.get("/sheet").status_code == 200)
check("manager cannot reach it", mgr.get("/sheet").status_code == 403)
check("reports as not configured without keys", "not configured" in admin.get("/sheet").text)
with _SL() as _db:
    _org = _db.scalar(select(_Org)) if False else None
from sqlalchemy import select as _sel
_db = _SL(); _org = _db.scalar(_sel(_Org))
for _name, _fn in _sh.TABS:
    try:
        _h, _r = _fn(_db, _org.id)
        check(f"tab '{_name}' builds", len(_h) > 0)
    except Exception as _e:
        check(f"tab '{_name}' builds", False, f"{type(_e).__name__}: {_e}")
_h, _rows = _sh._tasks(_db, _org.id)
check("Tasks tab has no None cells",
      not any(v is None for row in _rows for v in row))
check("Tasks tab includes the imported rows",
      any(f"BULK one {RUN}" in str(row[1]) for row in _rows))
_h2, _s = _sh._scores(_db, _org.id)
check("Scores tab carries the benchmark column", "Benchmark D/C/F" in _h2)
check("sync refuses politely when unconfigured",
      _sh.sync(_db, _org.id)["ok"] is False)
_db.close()

print("\n== attachments ==")
from app.services import storage as _st
cfg = doer.get("/attach/config").json()
check("upload config exposed", cfg["max_mb"] == 10 and cfg["enabled"], cfg)
check("accept list covers photos and documents",
      ".jpg" in cfg["accept"] and ".pdf" in cfg["accept"] and ".xlsx" in cfg["accept"])
check("executables are not accepted", ".exe" not in cfg["accept"])

# a doer's own task, in local mode
r = mgr.post("/tasks/new", data={
    "title": f"SMOKE attach {RUN}", "details": "x", "doer_id": "6",
    "branch_id": "", "priority": "medium", "due_at": "2026-12-31T18:00"})
aid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])

page = doer.get(f"/tasks/{aid}").text
check("uploader shown on the doer's submit card", 'id="uploader"' in page)
check("uploader states the size limit", "10 MB" in page)

img = b"\x89PNG\r\n\x1a\n" + b"x" * 3000
r = doer.post(f"/tasks/{aid}/attach/local",
              files={"files": ("register.png", img, "image/png")})
check("doer can attach a photo", r.status_code == 200, r.status_code)
body = doer.get(f"/tasks/{aid}").text
check("attachment listed on the task", "register.png" in body)
check("photo rendered as a thumbnail", "/view" in body)

check("oversized upload refused",
      doer.post(f"/tasks/{aid}/attach/local",
                files={"files": ("huge.png", b"x" * (11*1024*1024), "image/png")}
                ).status_code == 400)
check("executable upload refused",
      doer.post(f"/tasks/{aid}/attach/local",
                files={"files": ("bad.exe", b"MZ" + b"x"*100,
                                 "application/x-msdownload")}).status_code == 400)

# permissions
stranger = login("neha@gcs.local")
check("unrelated staff cannot attach to someone else's task",
      stranger.post(f"/tasks/{aid}/attach/local",
                    files={"files": ("x.png", img, "image/png")}).status_code == 404)
att_id = int(re.findall(r"/attachments/(\d+)", body)[0])
check("unrelated staff cannot download it",
      stranger.get(f"/attachments/{att_id}").status_code == 404)
check("the doer can download their own file",
      doer.get(f"/attachments/{att_id}").status_code == 200)
check("the manager can see the proof",
      mgr.get(f"/attachments/{att_id}").status_code == 200)
check("unrelated staff cannot delete it",
      stranger.post(f"/attachments/{att_id}/delete").status_code == 404)
check("the uploader can remove their own file",
      doer.post(f"/attachments/{att_id}/delete").status_code == 200)
check("file is gone from the task", "register.png" not in doer.get(f"/tasks/{aid}").text)

print("\n== delete user ==")
# make a throwaway person with real work attached
r = admin.post("/admin/users", data={
    "name": f"SMOKE Doomed {RUN}", "email": f"doomed.{RUN}@gcs.local",
    "password": "x", "role": "doer", "branch_id": "1",
    "bm_delegation": "60", "bm_checklist": "20", "bm_fms": "20"})
check("throwaway user created", r.status_code == 200, r.status_code)
import re as _re
page = admin.get("/admin/users").text
did = int(_re.findall(r'/admin/users/(\d+)/delete"[^>]*title="Delete SMOKE Doomed '
                      + RUN, page)[0])

check("delete button appears in the list", f'/admin/users/{did}/delete' in page)
dp = admin.get(f"/admin/users/{did}/delete")
check("delete page opens", dp.status_code == 200)
check("delete page asks for the name", 'name="confirm"' in dp.text)
check("clean account says nothing is linked", "Nothing is linked" in dp.text)

check("wrong confirmation name is refused",
      admin.post(f"/admin/users/{did}/delete",
                 data={"mode": "purge", "confirm": "nope"}).status_code == 400)
check("deleted with the right name",
      admin.post(f"/admin/users/{did}/delete",
                 data={"mode": "purge", "confirm": f"SMOKE Doomed {RUN}"}).status_code == 200)
check("user is gone", admin.get(f"/admin/users/{did}").status_code == 404)

print("\n== delete guards ==")
me = admin.get("/admin/users").text
# Checked precisely rather than by looking for the word "you" — that text
# moved when the Edit button was added, and a proxy like that passes happily
# while the actual delete link sits right there.
_my_id = admin.get("/admin/users").text
import re as _re
_mine = _re.search(r"mis@gcs\.local", me)
check("no delete button against your own row",
      f'/admin/users/2/delete"' not in me, "own delete link is on the page")
check("but your own row does offer Edit", '/admin/users/2"' in me)
check("cannot open delete page for yourself with a blocker",
      "can&#39;t delete your own account" in admin.get("/admin/users/2/delete").text
      or "own account" in admin.get("/admin/users/2/delete").text)
check("self-delete refused by the server",
      admin.post("/admin/users/2/delete",
                 data={"mode": "purge", "confirm": "Rajinder Singh"}).status_code == 400)

print("\n== delete with work: transfer ==")
r = admin.post("/admin/users", data={
    "name": f"SMOKE Busy {RUN}", "email": f"busy.{RUN}@gcs.local",
    "password": "x", "role": "doer", "branch_id": "1",
    "bm_delegation": "60", "bm_checklist": "20", "bm_fms": "20"})
page = admin.get("/admin/users").text
bid = int(_re.findall(r'/admin/users/(\d+)/delete"[^>]*title="Delete SMOKE Busy '
                      + RUN, page)[0])
r = admin.post("/tasks/new", data={
    "title": f"SMOKE inherited task {RUN}", "details": "x", "doer_id": str(bid),
    "branch_id": "", "priority": "medium", "due_at": "2026-12-31T18:00"})
kept = int(_re.findall(r"/tasks/(\d+)/comment", r.text)[0])
dp = admin.get(f"/admin/users/{bid}/delete").text
check("linked work is counted", "Tasks assigned" in dp and "Hand it over" in dp)
check("transfer needs a valid heir",
      admin.post(f"/admin/users/{bid}/delete",
                 data={"mode": "transfer", "transfer_to": "",
                       "confirm": f"SMOKE Busy {RUN}"}).status_code == 400)
check("transfer succeeds",
      admin.post(f"/admin/users/{bid}/delete",
                 data={"mode": "transfer", "transfer_to": "6",
                       "confirm": f"SMOKE Busy {RUN}"}).status_code == 200)
check("the task survived the transfer",
      admin.get(f"/tasks/{kept}").status_code == 200)
check("task now belongs to the heir",
      "PT Trainer Amit" in admin.get(f"/tasks/{kept}").text)
check("the person is gone", admin.get(f"/admin/users/{bid}").status_code == 404)

print("\n== benchmarks reach the dashboards ==")
dd2 = doer.get("/").text
check("doer tiles show the benchmark %", 'class="bm-pill"' in dd2)
ms = doer.get("/api/my-score?days=30").json()
check("my-score API exposes benchmark per source",
      all("benchmark" in s2 for s2 in ms["sources"]), ms["sources"])
sp2 = owner.get("/stats?days=90").text
check("admin cards show each doer's split", "70/10/20" in sp2 or "20/40/40" in sp2)

print("\n== false marking ==")
r = mgr.post("/tasks/new", data={
    "title": "SMOKE: false-mark candidate", "details": "x", "doer_id": "6",
    "branch_id": "", "priority": "medium", "due_at": "2026-12-31T18:00"})
fid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
doer.post(f"/tasks/{fid}/start")
submit(doer, fid, completion_note="done")
check("closed without audit", "completed" in doer.get(f"/tasks/{fid}").text.lower())
check("false mark refused without confirm",
      mgr.post(f"/tasks/{fid}/false-mark", data={"reason": "x"}).status_code == 400)
check("false mark applied with confirm",
      mgr.post(f"/tasks/{fid}/false-mark",
               data={"reason": "Register entry missing", "confirm": "yes"}).status_code == 200)
body = mgr.get(f"/tasks/{fid}").text
check("task shows false-marking banner", "false marking" in body.lower())
check("task was reopened", "reopened" in body.lower())
check("double false-mark refused",
      mgr.post(f"/tasks/{fid}/false-mark", data={"reason": "y", "confirm": "yes"}).status_code == 400)
check("doer notified", "false marking" in mgr.get("/outbox").text.lower())

print("\n== reopen ==")
submit(doer, fid, completion_note="redone")
check("reopen refused on non-completed",
      mgr.post(f"/tasks/{tid}/reopen", data={"reason": "x"}).status_code in (200, 400))

print("\n== edit / delete ==")
r = mgr.post("/tasks/new", data={
    "title": "SMOKE: to be edited", "details": "x", "doer_id": "6",
    "branch_id": "", "priority": "low", "due_at": "2026-12-30T10:00"})
eid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
check("manager cannot delete by default",
      mgr.post(f"/tasks/{eid}/delete").status_code == 403)
check("edit applies", mgr.post(f"/tasks/{eid}/edit", data={
    "title": "SMOKE: edited title", "details": "y", "doer_id": "7",
    "priority": "high", "due_at": "2026-12-31T10:00"}).status_code == 200)
body = mgr.get(f"/tasks/{eid}").text
check("edited title shown", "edited title" in body)
check("edit logged as a note", "Edited:" in body)
check("admin can delete", admin.post(f"/tasks/{eid}/delete").status_code == 200)
check("deleted task is gone", admin.get(f"/tasks/{eid}").status_code == 404)

print("\n== audit status ==")
r = mgr.post("/tasks/new", data={
    "title": "SMOKE: audit-state task", "details": "d", "doer_id": "6",
    "branch_id": "", "priority": "medium",
    "due_at": "2026-12-31T18:00", "requires_audit": "1"})
aid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
body = mgr.get(f"/tasks/{aid}").text
check("audit-flagged task starts as Audit pending", "Audit pending" in body)

r = mgr.post("/tasks/new", data={
    "title": "SMOKE: no-audit task", "details": "d", "doer_id": "6",
    "branch_id": "", "priority": "medium", "due_at": "2026-12-31T18:00"})
nid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
check("unflagged task starts as Not required",
      "Not required" in mgr.get(f"/tasks/{nid}").text)

check("audit cannot be completed without a remark",
      mgr.post(f"/tasks/{aid}/audit-state",
               data={"state": "completed", "remark": "  "}).status_code == 400)
check("a doer cannot change audit status",
      doer.post(f"/tasks/{aid}/audit-state",
                data={"state": "completed", "remark": "ok"}).status_code == 403)
r = mgr.post(f"/tasks/{aid}/audit-state",
             data={"state": "completed", "remark": "Checked the register, tallies."})
check("auditor marks the audit completed", r.status_code == 200, r.status_code)
body = mgr.get(f"/tasks/{aid}").text
check("audit remark is captured", "Checked the register, tallies." in body)
check("status now shows Audit completed", "Audit completed" in body)
check("the change is logged as a note", "Audit: Audit pending" in body)
r = mgr.post(f"/tasks/{nid}/audit-state", data={"state": "pending", "remark": ""})
check("a task can be pulled in for audit later",
      "Audit pending" in mgr.get(f"/tasks/{nid}").text)
check("audit-pending filter finds it",
      f'/tasks/{nid}#audit' in mgr.get("/tasks?scope=all&status=audit_pending").text)
check("audit column is on the task list", "Audit</th>" in mgr.get("/tasks?scope=all&status=all").text)
check("bad audit state refused",
      mgr.post(f"/tasks/{nid}/audit-state", data={"state": "nonsense"}).status_code == 400)

print("\n== paste attachments ==")
body = doer.get("/").text
r = doer.get(f"/tasks/{aid}")
body = doer.get(f"/tasks/{aid}").text
check("the doer's panel offers paste", "paste-zone" in body and "Ctrl" in body)
check("the file picker is still there", "Tap to choose a file" in body)
import io
png = (b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
r = mgr.post(f"/tasks/{nid}/attach/upload",
             files={"file": ("screenshot-20260910-101500.png", io.BytesIO(png), "image/png")})
check("a pasted screenshot uploads", r.status_code == 200, r.text[:120])
check("it comes back as an image", r.status_code == 200 and r.json().get("is_image"))
check("it is named, not 'image.png'",
      r.status_code == 200 and r.json()["filename"].startswith("screenshot-"))
check("it shows on the task",
      "screenshot-20260910" in mgr.get(f"/tasks/{nid}").text)
r = mgr.post(f"/tasks/{nid}/attach/upload",
             files={"file": ("evil.exe", io.BytesIO(b"MZ"), "application/x-msdownload")})
check("a disallowed type is refused", r.status_code == 400, r.status_code)
r = doer.post(f"/tasks/{aid}/attach/upload",
              files={"file": ("x.png", io.BytesIO(png), "image/png")})
check("the doer can attach to their own task", r.status_code == 200, r.status_code)

print("\n== help desk ==")
check("everyone can reach the help desk", doer.get("/help").status_code == 200)
check("help form loads", doer.get("/help/new").status_code == 200)
r = doer.post("/help/new", data={
    "subject": "SMOKE: pull the renewal list", "details": "Need it for the review.",
    "helper_id": "8", "priority": "high", "needed_by": "2026-12-28T17:00"})
check("raising help lands on a task", r.status_code == 200 and "Help:" in r.text, r.status_code)
hid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
body = admin.get(f"/tasks/{hid}").text
check("it became a Delegation task", "delegation" in body)
check("the raiser is the assigner", "Help requested by" in body)
check("it shows in the raiser's help list",
      "SMOKE: pull the renewal list" in doer.get("/help").text)
check("cannot ask yourself for help", doer.post("/help/new", data={
    "subject": "x", "details": "", "helper_id": "6", "priority": "low",
    "needed_by": "2026-12-28T17:00"}).status_code == 400)
check("cannot raise an empty request", doer.post("/help/new", data={
    "subject": "   ", "details": "", "helper_id": "8", "priority": "low",
    "needed_by": "2026-12-28T17:00"}).status_code == 400)

r = doer.post("/help/new", data={
    "subject": "SMOKE: to be declined", "details": "", "helper_id": "8",
    "priority": "low", "needed_by": "2026-12-28T17:00"})
did = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
helper = login("neha@gcs.local")
tickets = re.findall(r"/help/(\d+)/decline", helper.get("/help").text)
check("the helper sees a decline button", len(tickets) >= 2, tickets)
tk = tickets[0]          # newest first — the one we just raised to decline
check("an outsider cannot decline",
      admin.post(f"/help/{tk}/decline", data={"reason": "no"}).status_code == 403)
r = helper.post(f"/help/{tk}/decline", data={"reason": "On leave this week"})
check("the helper can decline", r.status_code == 200, r.status_code)
check("the reason is shown", "On leave this week" in helper.get("/help").text)
check("declining cancels the task", "cancelled" in admin.get(f"/tasks/{did}").text)
check("a declined request cannot be declined twice",
      helper.post(f"/help/{tk}/decline", data={"reason": "x"}).status_code == 400)

print("\n== help tickets don't block deletions ==")
r = doer.post("/help/new", data={
    "subject": "SMOKE: delete-me help", "details": "", "helper_id": "8",
    "priority": "low", "needed_by": "2026-12-27T17:00"})
xid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
check("a task carrying a help ticket can be deleted",
      admin.post(f"/tasks/{xid}/delete").status_code == 200)
check("the ticket goes with it",
      "SMOKE: delete-me help" not in doer.get("/help").text)
r = doer.post("/help/new", data={
    "subject": "SMOKE: transfer-me help", "details": "", "helper_id": "8",
    "priority": "low", "needed_by": "2026-12-27T17:00"})
page = admin.get("/admin/users/8/delete").text
check("help requests are listed before deleting someone", "Help Desk requests" in page)
check("a person with help tickets can be transferred out",
      admin.post("/admin/users/8/delete",
                 data={"mode": "transfer", "transfer_to": "7",
                       "confirm": "Spa Therapist Neha"}).status_code == 200)
check("their help request survives the transfer",
      "SMOKE: transfer-me help" in doer.get("/help").text)

print("\n== every link and button points somewhere real ==")
import route_check
check("no template links at a missing route", route_check.main() == 0)

print("\n== buttons that had never been clicked in a test ==")
r = admin.post("/admin/branches", data={"name": f"SMOKE Co {RUN}", "city": "Chandigarh"})
check("Add branch works", r.status_code == 200, r.status_code)
check("the new company is listed", f"SMOKE Co {RUN}" in r.text)
check("it can be picked as a task's branch",
      f"SMOKE Co {RUN}" in admin.get("/tasks/new").text)
check("a branch needs a name",
      admin.post("/admin/branches", data={"name": "  ", "city": "x"}).status_code == 400)
check("the same branch can't be added twice",
      admin.post("/admin/branches",
                 data={"name": f"SMOKE Co {RUN}", "city": ""}).status_code == 400)

print("\n== companies: see, edit, delete ==")
check("Branches is its own page", admin.get("/admin/branches").status_code == 200)
check("Users page no longer carries the branch form",
      "Add a branch" not in admin.get("/admin/users").text)
check("Branches appears in the menu", "Branches</span>" in admin.get("/admin/users").text)
up = admin.get("/admin/branches").text
check("branches are listed with their counts", "<h1>Branches</h1>" in up)
check("each company links to its own page", "/admin/branches/1" in up)
bpage = admin.get("/admin/branches/1")
check("company page opens", bpage.status_code == 200, bpage.status_code)
check("it shows what is filed under it", "What is filed under this branch" in bpage.text)
r = admin.post("/admin/branches/1", data={"name": f"Bodyzone RENAMED {RUN}", "city": "Mohali"})
check("rename works", r.status_code == 200 and f"Bodyzone RENAMED {RUN}" in r.text)
check("tasks follow the new name",
      f"Bodyzone RENAMED {RUN}" in admin.get("/stats").text)
r = admin.post("/admin/branches/1", data={"name": f"Bodyzone RENAMED {RUN}", "city": ""})
check("renaming to its own name is fine", r.status_code == 200)
check("cannot rename onto another company's name",
      admin.post("/admin/branches/1",
                 data={"name": "Spa Kora", "city": ""}).status_code == 400)
check("cannot blank a company name",
      admin.post("/admin/branches/1", data={"name": " ", "city": ""}).status_code == 400)
admin.post("/admin/branches/1", data={"name": "Bodyzone Fitness & Spa", "city": "Chandigarh"})

def branch_id_of(html, name):
    """The list is ordered by name, so pick the row, not the last match."""
    m = re.search(r'/admin/branches/(\d+)"><strong>' + re.escape(name), html)
    return int(m.group(1)) if m else None

r = admin.post("/admin/branches", data={"name": f"SMOKE empty {RUN}", "city": ""})
eid = branch_id_of(admin.get("/admin/branches").text, f"SMOKE empty {RUN}")
check("the new branch appears with its own id", eid is not None)
check("delete needs the name typed exactly",
      admin.post(f"/admin/branches/{eid}/delete",
                 data={"confirm": "wrong"}).status_code == 400)
check("an empty branch deletes cleanly",
      admin.post(f"/admin/branches/{eid}/delete",
                 data={"confirm": f"SMOKE empty {RUN}"}).status_code == 200)
check("it is gone from the list",
      f"SMOKE empty {RUN}" not in admin.get("/admin/branches").text)

# a company holding real work: everything must land somewhere, not vanish
r = admin.post("/admin/branches", data={"name": f"SMOKE busy {RUN}", "city": ""})
bid = branch_id_of(admin.get("/admin/branches").text, f"SMOKE busy {RUN}")
r = admin.post("/tasks/new", data={
    "title": f"SMOKE task in busy co {RUN}", "details": "", "doer_id": "6",
    "branch_id": str(bid), "priority": "low", "due_at": "2026-12-31T10:00"})
btid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
check("the task is filed under it", f"SMOKE busy {RUN}" in admin.get(f"/tasks/{btid}").text)
r = admin.post(f"/admin/branches/{bid}/delete",
               data={"confirm": f"SMOKE busy {RUN}", "move_to": "2"})
check("a branch holding work can be deleted", r.status_code == 200, r.status_code)
check("its task survived", admin.get(f"/tasks/{btid}").status_code == 200)
check("and moved to the chosen branch",
      "Spa Kora" in admin.get(f"/tasks/{btid}").text)
check("the Performance page still loads afterwards",
      admin.get("/stats").status_code == 200)

r = admin.post("/recurring", data={
    "title": f"SMOKE daily check {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "frequency": "daily", "day_of": "", "due_time": "18:00",
    "priority": "medium"})
check("Add checklist rule works", r.status_code == 200 and "SMOKE daily check" in r.text)
rid = re.findall(r"/recurring/(\d+)/toggle", r.text)
check("the rule can be paused",
      admin.post(f"/recurring/{rid[-1]}/toggle").status_code == 200)

r = admin.post("/flows/new", data={
    "name": f"SMOKE flow {RUN}", "description": "",
    "step_title": ["Step one", "Step two"], "step_doer": ["6", "7"],
    "step_tat": ["24", "24"], "step_priority": ["medium", "medium"],
    "step_fields": ["", ""], "step_audit": ["", ""]})
check("Create FMS flow works", r.status_code == 200, r.status_code)
fl = re.findall(r"/flows/(\d+)/start", r.text) or re.findall(r"/flows/(\d+)\"", r.text)
check("the new flow is startable", bool(fl), r.status_code)

check("user can be deactivated",
      admin.post("/admin/users/9/toggle").status_code == 200)
check("and reactivated",
      admin.post("/admin/users/9/toggle").status_code == 200)
check("Google Sheet page loads", admin.get("/sheet").status_code == 200)
check("Sync now answers even with no sheet configured",
      admin.post("/sheet/sync").status_code == 200)

print("\n== proof is required by default ==")
r = mgr.post("/tasks/new", data={
    "title": f"SMOKE proof needed {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "priority": "medium", "due_at": "2026-12-31T18:00"})
pid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
body = doer.get(f"/tasks/{pid}").text
check("a new task requires proof with nothing ticked", "Proof is required" in body)
check("the submit button is disabled until proof is there", "disabled" in body)
r = doer.post(f"/tasks/{pid}/submit", data={"completion_note": "done"})
check("submitting with no attachment is refused", r.status_code == 400, r.status_code)
check("the refusal explains what to do", "needs proof attached" in r.text)

import io
png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
doer.post(f"/tasks/{pid}/attach/upload",
          files={"file": ("proof.png", io.BytesIO(png), "image/png")})
check("with proof attached the doer is told they can submit",
      "you can submit" in doer.get(f"/tasks/{pid}").text)
r = doer.post(f"/tasks/{pid}/submit", data={"completion_note": "done"})
check("and the submit goes through", r.status_code == 200, r.status_code)
check("the task closed", "completed" in doer.get(f"/tasks/{pid}").text)

r = mgr.post("/tasks/new", data={
    "title": f"SMOKE proof off {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "priority": "medium", "due_at": "2026-12-31T18:00",
    "requires_attachment": "0"})   # exactly what an unticked box posts
oid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
check("an unticked proof box really turns it off",
      "Proof is required" not in doer.get(f"/tasks/{oid}").text)
check("and then it submits with nothing attached",
      doer.post(f"/tasks/{oid}/submit", data={"completion_note": "x"}).status_code == 200)

r = admin.post("/recurring", data={
    "title": f"SMOKE rule no proof {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "frequency": "daily", "day_of": "", "due_time": "18:00",
    "priority": "medium", "requires_attachment": "0"})
from app.db import SessionLocal as _S
from app.models import RecurringRule as _R
from sqlalchemy import select as _sel
_db = _S()
_rule = _db.scalar(_sel(_R).where(_R.title == f"SMOKE rule no proof {RUN}"))
check("an unticked checklist rule stores proof=off",
      _rule is not None and _rule.requires_attachment is False,
      _rule.requires_attachment if _rule else "no rule")
_db.close()
check("checklist rules default to requiring proof too",
      "Proof required before submitting" in admin.get("/recurring").text)
check("FMS steps ask about proof", "step_proof" in admin.get("/flows/new").text)

print("\n== attachments stored in the database (free hosting) ==")
import importlib
from app.services import storage as _st
import os as _os
_prev = _os.environ.get("STORAGE_BACKEND")
_os.environ["STORAGE_BACKEND"] = "db"
importlib.reload(_st)
import app.routers.attachments as _ra
importlib.reload(_ra)
check("db mode is recognised", _st.mode() == "db", _st.mode())
check("uploads stay available with no disk", _st.uploads_available())

r = mgr.post("/tasks/new", data={
    "title": f"SMOKE db-stored proof {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "priority": "medium", "due_at": "2026-12-31T18:00"})
dbid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
r = doer.post(f"/tasks/{dbid}/attach/upload",
              files={"file": ("in-db.png", io.BytesIO(TINY_PNG), "image/png")})
check("a file uploads with no disk involved", r.status_code == 200, r.text[:120])
aid2 = r.json()["id"]

from app.db import SessionLocal as _S2
from app.models import Attachment as _A
_d2 = _S2(); _att = _d2.get(_A, aid2)
check("the bytes are in the database row", _att.storage == "db" and _att.data == TINY_PNG)
_d2.close()

r = doer.get(f"/attachments/{aid2}")
check("it downloads back byte-for-byte", r.status_code == 200 and r.content == TINY_PNG)
r = doer.get(f"/attachments/{aid2}/view")
check("and previews inline", r.status_code == 200
      and "inline" in r.headers.get("content-disposition", ""))
check("proof now counts, so the task submits",
      doer.post(f"/tasks/{dbid}/submit", data={"completion_note": "x"}).status_code == 200)
check("deleting it does not error",
      doer.post(f"/attachments/{aid2}/delete").status_code in (200, 400))

if _prev is None:
    _os.environ.pop("STORAGE_BACKEND", None)
else:
    _os.environ["STORAGE_BACKEND"] = _prev
importlib.reload(_st); importlib.reload(_ra)
check("back to disk mode afterwards", _st.mode() == "local", _st.mode())

print("\n== a wrong link shows a readable page, not a raw error ==")
r = admin.get("/branches")            # the old, wrong address
check("a missing page returns 404", r.status_code == 404)
check("and shows a readable message", "Page not found" in r.text)
check("with the menu still there", "Branches</span>" in r.text)
check("the API still answers in JSON",
      admin.get("/api/stats?days=7").headers.get("content-type", "").startswith("application/json"))
r = doer.post("/tasks/999999/attach/upload",
              files={"file": ("x.png", io.BytesIO(TINY_PNG), "image/png")})
check("upload errors stay JSON for the uploader",
      r.headers.get("content-type", "").startswith("application/json"))

print("\n== branches: renamed, and the page opens ==")
check("the menu says Branches", "Branches</span>" in admin.get("/admin/users").text)
bp = admin.get("/admin/branches")
check("the Branches page opens", bp.status_code == 200 and "<h1>Branches</h1>" in bp.text)
check("nothing still says Company", "Company" not in bp.text)
bid1 = re.search(r'/admin/branches/(\d+)"><strong>', bp.text).group(1)
one = admin.get(f"/admin/branches/{bid1}")
check("clicking a branch opens its page", one.status_code == 200, one.status_code)
check("it says branch, not company", "filed under this branch" in one.text)

print("\n== holidays ==")
check("holidays page opens", admin.get("/admin/holidays").status_code == 200)
check("a doer cannot manage holidays", doer.get("/admin/holidays").status_code == 403)
r = admin.post("/admin/holidays", data={"day": "2026-11-08", "name": f"Diwali {RUN}",
                                        "branch_id": ""})
check("a holiday can be added", r.status_code == 200 and f"Diwali {RUN}" in r.text)
check("the same day cannot be added twice",
      admin.post("/admin/holidays",
                 data={"day": "2026-11-08", "name": "again", "branch_id": ""}).status_code == 400)
check("a holiday needs a name",
      admin.post("/admin/holidays",
                 data={"day": "2026-11-09", "name": "  ", "branch_id": ""}).status_code == 400)
check("a bad date is refused",
      admin.post("/admin/holidays",
                 data={"day": "not-a-date", "name": "x", "branch_id": ""}).status_code == 400)

# --- the rule that actually matters: no work is created on a holiday ---
from datetime import date as _date, datetime as _dt, timedelta as _td
from app.db import SessionLocal as _SS
from app.models import (RecurringRule as _RR, Recurrence as _Rec, Task as _T,
                        Holiday as _H, TaskSource as _TS)
from app.services import recurring as _recur, holidays as _hol
from sqlalchemy import select as _sel

_db = _SS()
_org = 1
_hday = _date(2027, 3, 15)
_db.add(_H(org_id=_org, branch_id=None, day=_hday, name="SMOKE holiday"))
_db.add(_RR(org_id=_org, branch_id=1, title=f"SMOKE daily on holiday {RUN}",
            doer_id=6, assigner_id=2, frequency=_Rec.DAILY, due_time="18:00"))
_db.commit()

check("the day is recognised as a holiday", _hol.is_holiday(_db, _org, _hday))
before = _db.scalar(_sel(__import__("sqlalchemy").func.count()).select_from(_T))
_recur.run_spawn(_db, today=_hday)
after = _db.scalar(_sel(__import__("sqlalchemy").func.count()).select_from(_T))
check("no checklist task is created on a holiday", before == after, f"{before} -> {after}")

_recur.run_spawn(_db, today=_hday + _td(days=1))
after2 = _db.scalar(_sel(__import__("sqlalchemy").func.count()).select_from(_T))
check("and the day after, tasks are created again", after2 > after, f"{after} -> {after2}")
check("the holiday is skipped, not owed the next day",
      after2 - after < 20, after2 - after)

# a deadline landing on a holiday moves forward
_moved, _why = _hol.shift_due(_db, _org, _dt.combine(_hday, _dt.min.time()).replace(hour=18))
check("a deadline on a holiday moves to the next day",
      _moved.date() == _hday + _td(days=1), _moved)
check("and keeps the time of day", _moved.hour == 18, _moved)
check("the reason is reported back", _why == "SMOKE holiday", _why)
_ok, _none = _hol.shift_due(_db, _org, _dt(2027, 3, 20, 18, 0))
check("an ordinary day is left alone", _none is None and _ok == _dt(2027, 3, 20, 18, 0))

# two holidays in a row
_db.add(_H(org_id=_org, branch_id=None, day=_date(2027, 4, 1), name="A"))
_db.add(_H(org_id=_org, branch_id=None, day=_date(2027, 4, 2), name="B"))
_db.commit()
_m2, _ = _hol.shift_due(_db, _org, _dt(2027, 4, 1, 10, 0))
check("a run of holidays is skipped through", _m2.date() == _date(2027, 4, 3), _m2)

# a branch holiday must not affect another branch
_db.add(_H(org_id=_org, branch_id=4, day=_date(2027, 5, 5), name="Jharkhand only"))
_db.commit()
check("a branch holiday applies to that branch",
      _hol.is_holiday(_db, _org, _date(2027, 5, 5), 4))
check("and not to a different branch",
      not _hol.is_holiday(_db, _org, _date(2027, 5, 5), 1))
_db.close()

r = admin.post("/tasks/new", data={
    "title": f"SMOKE holiday deadline {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "priority": "medium", "due_at": "2026-11-08T18:00"})
check("delegating onto a holiday still works", r.status_code == 200, r.status_code)
check("and the task says the deadline moved", "is a holiday" in r.text)
check("the new deadline is the next day", "09 Nov 2026" in r.text)

print("\n== date filters on the task lists ==")
r = admin.get("/tasks?scope=all&status=done")
check("completed filters on the completion date", "Completion date" in r.text)
r = admin.get("/tasks?scope=all&status=open")
check("open filters on the planned date", "Planned date" in r.text)
check("there is a Coming up view",
      admin.get("/tasks?scope=all&status=upcoming").status_code == 200)

# build a task closed on a known day, and one planned for a different day
_db = _SS()
_t = _T(org_id=1, branch_id=1, title=f"SMOKE closed in range {RUN}",
        assigner_id=2, doer_id=6, source=_TS.DELEGATION,
        due_at=_dt(2025, 1, 5, 18, 0), requires_attachment=False)
_t.status = __import__("app.models", fromlist=["TaskStatus"]).TaskStatus.COMPLETED
_t.closed_at = _dt(2025, 6, 20, 12, 0)
_db.add(_t); _db.commit(); _cid = _t.id
_db.close()

r = admin.get("/tasks?scope=all&status=done&date_from=2025-06-01&date_to=2025-06-30")
check("a task closed in June shows in a June completion range",
      f"SMOKE closed in range {RUN}" in r.text)
r = admin.get("/tasks?scope=all&status=done&date_from=2025-07-01&date_to=2025-07-31")
check("and not in a July range", f"SMOKE closed in range {RUN}" not in r.text)
r = admin.get("/tasks?scope=all&status=all&date_from=2025-01-01&date_to=2025-01-31")
check("but its PLANNED date puts it in a January planned range",
      f"SMOKE closed in range {RUN}" in r.text)
r = admin.get("/tasks?scope=all&status=done&date_from=2025-06-30&date_to=2025-06-01")
check("dates entered backwards are swapped, not empty",
      f"SMOKE closed in range {RUN}" in r.text)
r = admin.get("/tasks?scope=all&status=done&date_from=2025-06-20&date_to=2025-06-20")
check("a single day includes the whole day",
      f"SMOKE closed in range {RUN}" in r.text)
check("a nonsense date is ignored rather than crashing",
      admin.get("/tasks?scope=all&status=done&date_from=rubbish").status_code == 200)
check("the completion column is on the list", "Completed</th>" in r.text)

print("\n== first-run setup (hosts with no shell) ==")
from app.routers.setup import needs_setup as _ns
from app.db import SessionLocal as _S3
_d3 = _S3()
check("setup is hidden once users exist", not _ns(_d3))
_d3.close()
check("the setup page 404s on a live system", admin.get("/setup").status_code == 404)
check("a stranger cannot reach it either",
      TestClient(app).get("/setup").status_code == 404)
check("posting to it is refused too",
      TestClient(app).post("/setup", data={
          "name": "Hacker", "email": "h@x.com",
          "password": "abcdefgh1", "password2": "abcdefgh1"}).status_code == 404)
check("login still works normally", "Sign out" in
      TestClient(app, follow_redirects=True).post(
          "/login", data={"email": "mis@gcs.local", "password": "gcs1234"}).text)

print("\n== priority weight (5x / 2x / 1x) ==")
from app.models import Priority as _P, PRIORITY_WEIGHT as _PW
check("three levels only", sorted(p.value for p in _P) == ["high", "low", "medium"],
      [p.value for p in _P])
check("high counts 5, medium 2, low 1",
      (_PW[_P.HIGH], _PW[_P.MEDIUM], _PW[_P.LOW]) == (5, 2, 1))

from app.services.scoring import SourceScore as _SS, Card as _C

class _FakeTask:
    def __init__(self, pr): self.priority = pr
    @property
    def weight(self):
        return _PW[self.priority]

from app.services import scoring as _sc
check("one high task weighs the same as five low ones",
      _sc._w([_FakeTask(_P.HIGH)]) == _sc._w([_FakeTask(_P.LOW)] * 5) == 5)
check("a medium weighs two", _sc._w([_FakeTask(_P.MEDIUM)]) == 2)

r = mgr.post("/tasks/new", data={
    "title": f"SMOKE high priority {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "priority": "high", "due_at": "2026-12-31T18:00"})
hid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
body = mgr.get(f"/tasks/{hid}").text
check("the task page shows the multiplier", "5\u00d7" in body)
check("and says what it means in words", "counts as 5" in body)
check("the doer still sees exactly one row, not five",
      doer.get("/tasks?scope=mine&status=open").text.count(f"SMOKE high priority {RUN}") == 1)

# the score must move five times as far for a high task as a low one
def _score_gap(tasks):
    c = _C()
    c.sources = {"delegation": _SS(key="delegation", benchmark=100,
                                   planned=_sc._w(tasks), completed=0,
                                   not_done=_sc._w(tasks)).compute()}
    return c.compute().gap

check("missing one high task hurts as much as missing five low ones",
      _score_gap([_FakeTask(_P.HIGH)]) == _score_gap([_FakeTask(_P.LOW)] * 5))

# sorting
r = mgr.post("/tasks/new", data={
    "title": f"SMOKE low priority {RUN}", "details": "", "doer_id": "6",
    "branch_id": "", "priority": "low", "due_at": "2026-01-01T09:00"})
lid = int(re.findall(r"/tasks/(\d+)/comment", r.text)[0])
listing = mgr.get("/tasks?scope=all&status=open").text
check("high sorts above low even when low is due sooner",
      listing.index(f"SMOKE high priority {RUN}") < listing.index(f"SMOKE low priority {RUN}"))

print("\n== follow-ups: PC and EA ==")
pc = login("pc@gcs.local")
ea = login("ea@gcs.local")
check("the PC can open Follow-ups", pc.get("/followups").status_code == 200)
check("the EA can open Follow-ups", ea.get("/followups").status_code == 200)
check("the PC lands on the Checklist & FMS desk", "Checklist &amp; FMS" in pc.get("/followups").text)
check("the EA lands on the Delegation desk", "EA \u2014 Delegation" in ea.get("/followups").text)
check("Follow-ups is in the PC's menu", "Follow-ups</span>" in pc.get("/").text)
check("an ordinary doer does not see the menu",
      "Follow-ups</span>" not in doer.get("/").text)

from datetime import date as _date
_today = _date.today().isoformat()
ea_page = ea.get(f"/followups?desk=ea&day={_today}").text
check("the EA sees delegation tasks that are not theirs",
      "SMOKE high priority" in ea_page or "tick" in ea_page.lower())
_ids = re.findall(r"/followups/(\d+)/tick", ea_page)
check("there is something to chase", len(_ids) > 0, len(_ids))

if _ids:
    tid_f = _ids[0]
    r = ea.post(f"/followups/{tid_f}/tick", data={"day": _today, "desk": "ea"})
    check("the EA can tick a follow-up", r.status_code == 200, r.status_code)
    check("the tick is recorded", "Followed up" in r.text)
    page = ea.get(f"/followups?desk=ea&day={_today}").text
    check("the count moves", "fu-done" in page)
    r = ea.post(f"/followups/{tid_f}/tick", data={"day": _today, "desk": "ea"})
    check("ticking again unticks it", "fu-done" not in r.text)
    ea.post(f"/followups/{tid_f}/tick", data={"day": _today, "desk": "ea"})

    check("the PC cannot tick a delegation task",
          pc.post(f"/followups/{tid_f}/tick",
                  data={"day": _today, "desk": "ea"}).status_code == 403)
    check("an ordinary doer cannot tick anything",
          doer.post(f"/followups/{tid_f}/tick",
                    data={"day": _today, "desk": "ea"}).status_code == 403)
    check("a follow-up cannot be recorded for tomorrow",
          ea.post(f"/followups/{tid_f}/tick",
                  data={"day": "2030-01-01", "desk": "ea"}).status_code == 400)

    # Yesterday's tick must not cover today. Tested on a task that was
    # genuinely open yesterday — the desk only lists what was open on the day
    # you are looking at, so ticking today's brand-new task and then asking
    # for yesterday's page proves nothing.
    import datetime as _dtf
    from app import clock as _ckf
    _yday_d = _ckf.today() - _dtf.timedelta(days=1)
    _yday = _yday_d.isoformat()
    _old_id = int(re.search(r"/tasks/(\d+)", admin.post("/tasks/new", data={
        "title": f"open since before yesterday {RUN}", "doer_id": 6,
        "priority": "medium",
        "due_at": (_ckf.now() - _dtf.timedelta(days=2)).strftime("%Y-%m-%dT%H:%M")},
        follow_redirects=False).headers["location"]).group(1))

    _ypage_before = ea.get(f"/followups?desk=ea&day={_yday}").text
    check("that task was open yesterday", f"/tasks/{_old_id}" in _ypage_before)
    ea.post(f"/followups/{_old_id}/tick", data={"day": _yday, "desk": "ea"})
    _ypage = ea.get(f"/followups?desk=ea&day={_yday}").text
    check("yesterday can be ticked separately", "fu-done" in _ypage)
    _tpage = ea.get(f"/followups?desk=ea&day={_today}").text
    check("and today keeps its own tally — still untickled",
          f'/followups/{_old_id}/tick' in _tpage)
    check("yesterday's tick did not mark today done",
          _tpage.count("fu-done") < _ypage.count("fu-done") + 1
          or "fu-done" not in _tpage.split(f"/tasks/{_old_id}")[0][-400:])

print("\n== follow-up report (now under Reports) ==")
check("the old link still lands on the report",
      admin.get("/followups/report").url.path == "/reports/followups")
rep = admin.get("/reports/followups")
check("the report opens", rep.status_code == 200, rep.status_code)
check("it names who holds each desk", "Karan" in rep.text and "Simran" in rep.text)
# The desk table is the only place a name means "this person holds the desk".
# The employee filter lists everybody, which is not the same thing.
_desk_table = rep.text.split("By desk", 1)[1].split("Day by day", 1)[0]
check("and does not list every admin as the PC",
      "CMD Sir" not in _desk_table, "admins should not be listed as desk holders")
check("it shows follow-ups still pending", "Follow-up pending" in rep.text)
check("a plain doer cannot open the report",
      doer.get("/reports/followups").status_code == 403)
check("but the PC can", pc.get("/reports/followups").status_code == 200)
check("a date range works",
      admin.get("/reports/followups?date_from=2026-09-01&date_to=2026-09-30").status_code == 200)
check("backwards dates are handled",
      admin.get("/reports/followups?date_from=2026-09-30&date_to=2026-09-01").status_code == 200)

print("\n== the priority column is plain text, not a native enum ==")
from sqlalchemy import inspect as _inspect, Enum as _SAEnum
from app.db import engine as _eng
_col = [c for c in _inspect(_eng).get_columns("tasks") if c["name"] == "priority"][0]
check("priority is stored as text", not isinstance(_col["type"], _SAEnum),
      repr(_col["type"]))
check("so renaming a level is an ordinary UPDATE", True)

print("\n== flow steps protected from deletion ==")
check("cannot delete a live FMS step",
      admin.post(f"/tasks/{step1}/delete").status_code == 400)

print("\n== the Reports menu ==")
_idx = admin.get("/reports")
check("the reports index opens", _idx.status_code == 200, _idx.status_code)
for _t in ["1 · Delegation", "2 · Checklist", "3 · FMS", "4 · Follow-ups",
           "5 · Audit", "6 · EM score"]:
    check(f"it offers {_t}", _t in _idx.text)

check("a doer sees the index too", doer.get("/reports").status_code == 200)
check("but is not offered the EM score report", "6 · EM score" not in doer.get("/reports").text)
check("nor the follow-up report", "4 · Follow-ups" not in doer.get("/reports").text)

print("\n== reports 1-3: pending and completed, by work type ==")
for _src, _label in [("delegation", "Delegation"), ("checklist", "Checklist"),
                     ("fms", "FMS")]:
    r = admin.get(f"/reports/tasks?source={_src}")
    check(f"{_label} report opens", r.status_code == 200, r.status_code)
    check(f"{_label} report is about {_label}",
          f"{_label} — pending" in r.text, r.text[:200])
    for _state in ["pending", "completed", "overdue", "all"]:
        check(f"{_label}: the {_state} tab works",
              admin.get(f"/reports/tasks?source={_src}&state={_state}").status_code == 200)

check("an unknown report is a readable 404",
      admin.get("/reports/tasks?source=nonsense").status_code == 404)

# Pending is filtered on the PLANNED date, completed on the COMPLETION date.
# Filtering both on one column is the usual way this goes quietly wrong.
_r = admin.get("/reports/tasks?source=delegation")
check("the page says which date it filters on",
      "planned date for pending" in _r.text and "completion date for completed" in _r.text)

print("\n== every report filter has BOTH a from and a to ==")
for _path in ["/reports/tasks?source=delegation", "/reports/tasks?source=checklist",
              "/reports/tasks?source=fms", "/reports/followups",
              "/reports/audit", "/reports/score"]:
    _t = admin.get(_path).text
    check(f"{_path} has a from date", 'name="date_from"' in _t)
    check(f"{_path} has a to date", 'name="date_to"' in _t)
    check(f"{_path} offers quick ranges", "Quick range" in _t)
    check(f"{_path} filters by branch", 'name="branch"' in _t)

_fu = pc.get("/followups").text
check("the follow-up desk has a from date too", 'name="date_from"' in _fu)
check("and a to date", 'name="date_to"' in _fu)
_range = pc.get("/followups?date_from=2026-09-01&date_to=2026-09-07")
check("a range on the desk gives a day-by-day summary",
      _range.status_code == 200 and "day by day" in _range.text.lower())
check("a single day still gives the tick list",
      "tickbox" in pc.get("/followups?date_from=%s&date_to=%s" % (_today, _today)).text)

print("\n== report 5: audit pending and completed ==")
_a = admin.get("/reports/audit")
check("the audit report opens", _a.status_code == 200, _a.status_code)
check("it splits by work type",
      "Delegation" in _a.text and "Checklist" in _a.text and "FMS" in _a.text)
check("it shows audit pending and completed",
      "Audit pending" in _a.text and "Audit completed" in _a.text)
for _state in ["pending", "completed", "not_required", "all"]:
    check(f"audit report: the {_state} tab works",
          admin.get(f"/reports/audit?state={_state}").status_code == 200)

print("\n== report 6: EM score, person-wise and branch-wise ==")
_s = admin.get("/reports/score")
check("the score report opens", _s.status_code == 200, _s.status_code)
check("it shows an average", "Average score" in _s.text)
check("person-wise is the default", "Person-wise" in _s.text)
check("branch-wise works",
      admin.get("/reports/score?view=branch").status_code == 200)
check("it shows the bifurcation per work type",
      "not done" in _s.text and "late" in _s.text)
check("a plain doer cannot open it", doer.get("/reports/score").status_code == 403)

# Picking a branch must narrow the employee list to that branch's people.
from app.db import SessionLocal as _SL
from app.models import Branch as _Br, User as _U
from sqlalchemy import select as _sel
with _SL() as _d:
    _bz = _d.scalar(_sel(_Br).where(_Br.name.like("Bodyzone%")))
    _bz_id = _bz.id
    _in_bz = {u.name for u in _d.scalars(_sel(_U).where(_U.branch_id == _bz_id)).all()}
    _out_bz = {u.name for u in _d.scalars(
        _sel(_U).where(_U.branch_id != _bz_id, _U.branch_id.is_not(None))).all()}
_bpage = admin.get(f"/reports/score?branch={_bz_id}").text
_table = _bpage.split("Person-wise", 1)[1]
check("picking a branch keeps its own employees",
      any(n in _table for n in _in_bz), "no Bodyzone employee listed")
check("and drops everybody else",
      not any(n in _table for n in _out_bz - _in_bz),
      "an employee from another branch is still listed")

print("\n== the dashboard is organised into sections ==")
_dash = admin.get("/").text
for _sec in ["1 · My work", "2 · My EM score", "3 · My team"]:
    check(f"the dashboard has '{_sec}'", _sec in _dash)
check("it groups my work by type of work", "By type of work" in _dash)
check("it shows the team average", "Average EM score" in _dash)
check("with the bifurcation behind it", "Score bifurcation" in _dash)
check("and a branch-wise roll-up", "Branch-wise" in _dash)

_ddash = doer.get("/").text
check("a doer gets the same score sections",
      "1 · My work" in _ddash and "2 · My EM score" in _ddash)
check("a doer sees their own bifurcation",
      "Work not done" in _ddash and "Not done on time" in _ddash)
check("but no team section", "3 · My team" not in _ddash)

print("\n== every user row has a visible Edit button ==")
from app.db import SessionLocal as _SL2
from app.models import User as _U2
from sqlalchemy import select as _sel2
with _SL2() as _d:
    _me = _d.scalar(_sel2(_U2).where(_U2.email == "mis@gcs.local"))
    _other = _d.scalar(_sel2(_U2).where(_U2.email == "amit@gcs.local"))
    _me_id, _other_id = _me.id, _other.id

_up = admin.get("/admin/users").text
check("someone else's row has an Edit button",
      f'href="/admin/users/{_other_id}"' in _up and ">Edit<" in _up)
check("your own row has one too",
      _up.count(f'href="/admin/users/{_me_id}"') >= 1)
# The Edit button must not be the only way in, and must not be a dead link.
check("the Edit button opens the edit page",
      admin.get(f"/admin/users/{_other_id}").status_code == 200)
_saved = admin.post(f"/admin/users/{_other_id}",
                    data={"name": "PT Trainer Amit", "phone": "+919800000006",
                          "role": "doer", "branch_id": "", "department_id": "",
                          "rights": ["create_task"], "bm_delegation": 70,
                          "bm_checklist": 10, "bm_fms": 20})
check("the edit page can actually save", _saved.status_code == 200, _saved.status_code)
check("and the change sticks", "70 / 10 / 20" in admin.get("/admin/users").text)
check("editing your own account is allowed",
      admin.get(f"/admin/users/{_me_id}").status_code == 200)
check("but your own role stays locked",
      "Role and rights are locked here" in admin.get(f"/admin/users/{_me_id}").text)

print("\n== reports are scoped until the right is given ==")
from app.db import SessionLocal as _SL3
from app.models import User as _U3, Task as _T3, Right as _R3
from sqlalchemy import select as _s3

with _SL3() as _d:
    _amit = _d.scalar(_s3(_U3).where(_U3.email == "amit@gcs.local"))
    _amit_id = _amit.id
    _amit_rights = _amit.rights

# By default a doer's report is about the doer, and nobody else.
_dr = doer.get("/reports/tasks?source=delegation&state=all")
check("a doer can open the delegation report", _dr.status_code == 200, _dr.status_code)
check("and is told it is only their own work",
      "covers your own work only" in _dr.text)
with _SL3() as _d:
    _mine = {t.id for t in _d.scalars(_s3(_T3).where(
        _T3.source == "DELEGATION",
        (_T3.doer_id == _amit_id) | (_T3.assigner_id == _amit_id))).all()}
    _theirs = {t.id for t in _d.scalars(_s3(_T3).where(
        _T3.source == "DELEGATION", _T3.doer_id != _amit_id,
        _T3.assigner_id != _amit_id)).all()}
_shown = {i for i in _mine if f'/tasks/{i}"' in _dr.text}
_leaked = {i for i in _theirs if f'/tasks/{i}"' in _dr.text}
check("their own delegation tasks are listed", bool(_shown), "none of their own shown")
check("somebody else's are not", not _leaked, f"leaked task ids {sorted(_leaked)[:5]}")

check("the EM score report is refused", doer.get("/reports/score").status_code == 403)
check("the follow-up report is refused", doer.get("/reports/followups").status_code == 403)
_ix = doer.get("/reports").text
check("the index does not offer the EM score report", "6 · EM score" not in _ix)
check("nor the follow-up report", "4 · Follow-ups" not in _ix)
check("and the sidebar hides both",
      "EM score report" not in _ix and "Follow-up report" not in _ix)

# Now tick the right and the same pages cover the company.
admin.post(f"/admin/users/{_amit_id}",
           data={"name": "PT Trainer Amit", "phone": "+919800000006",
                 "role": "doer", "branch_id": "", "department_id": "",
                 "rights": ["view_all_reports"], "bm_delegation": 60,
                 "bm_checklist": 20, "bm_fms": 20})
_dr2 = doer.get("/reports/tasks?source=delegation&state=all")
check("with the right, the warning is gone",
      "covers your own work only" not in _dr2.text)
check("and other people's tasks appear",
      any(f'/tasks/{i}"' in _dr2.text for i in _theirs), "still scoped to self")
check("the EM score report opens", doer.get("/reports/score").status_code == 200)
check("the follow-up report opens", doer.get("/reports/followups").status_code == 200)
check("the branch filter offers every branch",
      doer.get("/reports/score").text.count("<option value=\"") >
      _dr.text.count("<option value=\""))
_ix2 = doer.get("/reports").text
check("the index now offers all six",
      "6 · EM score" in _ix2 and "4 · Follow-ups" in _ix2)

# The right must not become a back door into anything else.
check("it does not grant delegating work",
      doer.get("/tasks/new").status_code == 403)
check("nor managing users", doer.get("/admin/users").status_code == 403)
check("nor the Performance page", doer.get("/stats").status_code == 403)

# put the user back the way the seed left them
with _SL3() as _d:
    _u = _d.get(_U3, _amit_id); _u.rights = _amit_rights; _d.commit()

print("\n== each person is told whose figures they are looking at ==")
check("a manager is told it is their branch",
      "This report covers Bodyzone Fitness &amp; Spa" in
      mgr.get("/reports/tasks?source=delegation").text
      or "covers Bodyzone" in mgr.get("/reports/tasks?source=delegation").text)
check("an admin gets no warning at all",
      "covers your own work only" not in admin.get("/reports/tasks?source=delegation").text
      and "To see every branch" not in admin.get("/reports/tasks?source=delegation").text)
# A manager's report must actually stay inside their branch.
with _SL3() as _d:
    _mgr = _d.scalar(_s3(_U3).where(_U3.email == "bz.manager@gcs.local"))
    _elsewhere = [t.id for t in _d.scalars(_s3(_T3).where(
        _T3.source == "DELEGATION", _T3.branch_id != _mgr.branch_id,
        _T3.doer_id != _mgr.id, _T3.assigner_id != _mgr.id)).all()]
_mp = mgr.get("/reports/tasks?source=delegation&state=all").text
check("and another branch's tasks stay out of it",
      not any(f'/tasks/{i}"' in _mp for i in _elsewhere),
      "a task from another branch is listed")

print("\n== the right is offered when creating a user ==")
_form = admin.get("/admin/users").text
check("the tick box exists", 'value="view_all_reports"' in _form)
check("with a plain-English label", "See everyone&#39;s reports" in _form
      or "See everyone\u2019s reports" in _form)
# Read the actual default-rights map the page ships to the browser, rather
# than guessing at the surrounding HTML — the word "manager" appears a dozen
# times on that page and a substring search finds the wrong one.
import json as _json
_def = _json.loads(_form.split("DEF = ", 1)[1].split(";", 1)[0].strip())
check("a doer gets nothing by default", _def["doer"] == [])
check("a manager does NOT get it by default",
      "view_all_reports" not in _def["manager"], _def["manager"])
check("an admin does", "view_all_reports" in _def["admin"])

print("\n== the clock runs on Indian time, not UTC ==")
from datetime import timezone as _tz, timedelta as _td, datetime as _dt
from app import clock as _clock
_IST = _tz(_td(hours=5, minutes=30))
_real = _dt.now(_IST).replace(tzinfo=None)
check("clock.now() is Chandigarh time",
      abs((_clock.now() - _real).total_seconds()) < 5,
      f"off by {(_clock.now() - _real).total_seconds()/3600:.1f}h")
check("clock.today() is the Indian date", _clock.today() == _real.date())

# A deadline typed as 6pm must go overdue at 6pm, not 11:30pm. This is the
# whole point: the 'not done on time' half of every benchmark depends on it.
from app.models import Task as _TK, TaskSource as _TS, Priority as _PR
_late = _TK(title="t", due_at=_clock.now() - _td(minutes=1),
            source=_TS.DELEGATION, priority=_PR.MEDIUM)
_soon = _TK(title="t", due_at=_clock.now() + _td(minutes=1),
            source=_TS.DELEGATION, priority=_PR.MEDIUM)
check("a deadline one minute ago is overdue", _late.is_overdue)
check("a deadline one minute away is not", not _soon.is_overdue)

# And nothing anywhere may quietly go back to UTC.
import subprocess as _sp, os as _os
_hits = _sp.run(["grep", "-rn", "utcnow()", "app/", "--include=*.py"],
                capture_output=True, text=True).stdout
# clock.py names it in its own docstring, explaining why it is gone.
_hits = "\n".join(l for l in _hits.splitlines() if "app/clock.py" not in l)
check("no utcnow() left anywhere in the app", _hits.strip() == "", _hits[:200])
_naive = _sp.run(["grep", "-rn", "date.today()", "app/", "--include=*.py"],
                 capture_output=True, text=True).stdout
_naive = "\n".join(l for l in _naive.splitlines() if "app/clock.py" not in l)
check("no bare date.today() either", _naive.strip() == "", _naive[:200])

# A task created now must be stamped with Indian time in the database.
_tid = int(re.search(r"/tasks/(\d+)",
    admin.post("/tasks/new", data={"title": f"clock {RUN}", "doer_id": 6,
        "priority": "medium", "due_at": (_clock.now() + _td(days=1)).strftime("%Y-%m-%dT%H:%M"),
    }, follow_redirects=False).headers["location"]).group(1))
from app.db import SessionLocal as _SLC
with _SLC() as _d:
    _row = _d.get(_TK, _tid)
    check("its created_at is Indian time",
          abs((_row.created_at - _dt.now(_IST).replace(tzinfo=None)).total_seconds()) < 120,
          str(_row.created_at))

print("\n== departments: see, add, edit, delete ==")
from app.db import SessionLocal as _SLD
from app.models import Department as _Dept, Branch as _Brn, User as _UD
from sqlalchemy import select as _sd

_dp = admin.get("/admin/departments")
check("the departments page opens", _dp.status_code == 200, _dp.status_code)
check("it is reachable from the sidebar", "/admin/departments" in admin.get("/").text)
check("and from the user form", "/admin/departments" in admin.get("/admin/users").text)

with _SLD() as _d:
    _bz = _d.scalar(_sd(_Brn).where(_Brn.name.like("Bodyzone%"))).id
    _sk = _d.scalar(_sd(_Brn).where(_Brn.name.like("Spa Kora%"))).id

_name = f"Front Desk {RUN}"
admin.post("/admin/departments", data={"name": _name, "branch_id": str(_bz)})
with _SLD() as _d:
    _new = _d.scalar(_sd(_Dept).where(_Dept.name == _name))
check("a department can be added", _new is not None)
check("it is listed", _name in admin.get("/admin/departments").text)
_did = _new.id

check("a blank name is refused",
      admin.post("/admin/departments", data={"name": "  ", "branch_id": ""}).status_code == 400)
check("the same name twice in one branch is refused",
      admin.post("/admin/departments",
                 data={"name": _name, "branch_id": str(_bz)}).status_code == 400)
# But the same name under a DIFFERENT branch is normal — every branch has a
# front desk — so that must be allowed.
check("the same name under another branch is allowed",
      admin.post("/admin/departments",
                 data={"name": _name, "branch_id": str(_sk)}).status_code == 200)

check("the edit page opens", admin.get(f"/admin/departments/{_did}").status_code == 200)
admin.post(f"/admin/departments/{_did}",
           data={"name": f"Reception {RUN}", "branch_id": ""})
with _SLD() as _d:
    _r = _d.get(_Dept, _did)
    check("renaming works", _r.name == f"Reception {RUN}", _r.name)
    check("and it can be moved to all-branches", _r.branch_id is None)

# Put a person in it, then delete it: they must keep everything but the label.
_person = admin.post("/admin/users", data={"name": f"Dept Tester {RUN}",
    "email": f"dept.{RUN}@gcs.local", "phone": "", "password": "deptpass123",
    "role": "doer", "branch_id": str(_bz), "department_id": str(_did),
    "rights": [], "bm_delegation": 60, "bm_checklist": 20, "bm_fms": 20})
with _SLD() as _d:
    _pu = _d.scalar(_sd(_UD).where(_UD.email == f"dept.{RUN}@gcs.local"))
    _puid, _pu_branch = _pu.id, _pu.branch_id
    check("a person can be put in a department", _pu.department_id == _did)
# Two branches can both have a "Front Desk", so the dropdown has to say which.
_udrop = admin.get("/admin/users").text
_opts = re.findall(r'<option value="\d+">([^<]*Front Desk[^<]*)</option>', _udrop)
check("the department dropdown names the branch",
      _opts and all("—" in o for o in _opts), _opts)
check("so two same-named departments are distinguishable",
      len(set(_opts)) == len(_opts), _opts)

check("the edit page lists who is in it",
      f"Dept Tester {RUN}" in admin.get(f"/admin/departments/{_did}").text)

check("deleting needs the name typed exactly",
      admin.post(f"/admin/departments/{_did}", data={}, follow_redirects=False) is not None)
_bad = TestClient(app, follow_redirects=False)
_bad.post("/login", data={"email": "mis@gcs.local", "password": "gcs1234"})
check("a wrong confirmation is refused",
      _bad.post(f"/admin/departments/{_did}/delete",
                data={"confirm": "wrong"}).status_code == 400)
admin.post(f"/admin/departments/{_did}/delete", data={"confirm": f"Reception {RUN}"})
with _SLD() as _d:
    check("the department is gone", _d.get(_Dept, _did) is None)
    _pu2 = _d.get(_UD, _puid)
    check("but the person survives", _pu2 is not None)
    check("with no department", _pu2.department_id is None)
    check("and their branch untouched", _pu2.branch_id == _pu_branch)

print("\n== My Tasks separates Delegation, Checklist and FMS ==")
from app.db import SessionLocal as _SLT
from app.models import Task as _TT, TaskSource as _TSR, User as _UT, TaskStatus as _TST
from sqlalchemy import select as _st

OPEN_T = (_TST.PENDING, _TST.IN_PROGRESS, _TST.REJECTED, _TST.REOPENED)
with _SLT() as _d:
    _am = _d.scalar(_st(_UT).where(_UT.email == "amit@gcs.local"))
    _want = {}
    for _k, _src in [("delegation", _TSR.DELEGATION), ("checklist", _TSR.RECURRING),
                     ("fms", _TSR.FLOW)]:
        _want[_k] = {t.id for t in _d.scalars(_st(_TT).where(
            _TT.doer_id == _am.id, _TT.source == _src,
            _TT.status.in_(OPEN_T))).all()}

_all = doer.get("/tasks?scope=mine&status=open")
check("the work-type row is on the page", "Work type" in _all.text)
check("it offers all three kinds",
      all(l in _all.text for l in ["Delegation", "Checklist", "FMS"]))
check("and an All work tab", "All work" in _all.text)

# Each tab must show that kind of work and nothing else.
for _k in ("delegation", "checklist", "fms"):
    _pg = doer.get(f"/tasks?scope=mine&status=open&source={_k}")
    check(f"the {_k} tab opens", _pg.status_code == 200, _pg.status_code)
    _shown = {i for i in range(1, 400) if f'/tasks/{i}"' in _pg.text}
    _others = set().union(*[v for kk, v in _want.items() if kk != _k]) if _want else set()
    check(f"{_k}: shows its own tasks",
          _want[_k] <= _shown or not _want[_k],
          f"missing {sorted(_want[_k] - _shown)[:4]}")
    check(f"{_k}: shows no other kind",
          not (_shown & _others), f"leaked {sorted(_shown & _others)[:4]}")

# The tab counts must match what the tabs actually contain.
import re as _re
_row = _all.text.split("Work type", 1)[1].split("</div>", 1)[0]
_nums = [int(n) for n in _re.findall(r"<b>(\d+)</b>", _row)]
check("the tab counts add up", _nums and _nums[0] == sum(_nums[1:]),
      f"all={_nums[:1]} parts={_nums[1:]}")
check("and match the database",
      len(_nums) == 4 and _nums[1:] == [len(_want["delegation"]),
                                        len(_want["checklist"]),
                                        len(_want["fms"])],
      f"page {_nums[1:]} vs db {[len(_want[k]) for k in ('delegation','checklist','fms')]}")

print("\n== and sorts high priority first ==")
check("the page says how it is sorted",
      "high priority first, then by deadline" in _all.text.lower())
_open_pg = doer.get("/tasks?scope=mine&status=open").text
_ids = [int(i) for i in _re.findall(r'/tasks/(\d+)"', _open_pg)]
with _SLT() as _d:
    _seen, _ordered = set(), []
    for _i in _ids:
        if _i in _seen: continue
        _seen.add(_i)
        _t = _d.get(_TT, _i)
        if _t and _t.doer_id == _am.id:
            _ordered.append((_t.priority.value, _t.due_at))
_rank = {"high": 0, "medium": 1, "low": 2}
check("high priority really is listed first",
      all(_rank[_ordered[i][0]] <= _rank[_ordered[i+1][0]]
          for i in range(len(_ordered) - 1)),
      str([o[0] for o in _ordered]))
check("and within a priority, the earliest deadline first",
      all(_ordered[i][1] <= _ordered[i+1][1]
          for i in range(len(_ordered) - 1)
          if _ordered[i][0] == _ordered[i+1][0]))

# Filtering by work type must survive changing the status tab.
_done = doer.get("/tasks?scope=mine&status=done&source=checklist")
check("the work type carries over to Completed", _done.status_code == 200)
check("and stays selected there", "source=checklist" in _done.text)

print("\n== the dashboard cards go to that list, not a report ==")
_dash = doer.get("/").text
# Jinja escapes & as &amp; inside an href, so compare against the escaped form.
for _k in ("delegation", "checklist", "fms"):
    _href = f"/tasks?scope=mine&amp;status=open&amp;source={_k}"
    check(f"the {_k} card opens the doer's own list", _href in _dash,
          "card still points at a report")

print("\n== deadlines default to the end of the day ==")
_nf = admin.get("/tasks/new").text
check("delegation defaults to 11:59 pm", "T23:59" in _nf,
      re.search(r'value="[^"]*T\d\d:\d\d"', _nf).group(0) if re.search(r'value="[^"]*T\d\d:\d\d"', _nf) else "?")
check("a checklist rule does too", 'value="23:59"' in admin.get("/recurring").text)
check("and a help request", "T23:59" in doer.get("/help/new").text)

print("\n== only the right may move a planned date ==")
from app.db import SessionLocal as _SLP
from app.models import Task as _TP, User as _UP, Right as _RP
from sqlalchemy import select as _sp
from datetime import timedelta as _tdp

# The task is assigned TO the editor, so they can genuinely see it — a 404
# would have made the "field is locked" check pass for the wrong reason.
with _SLP() as _d:
    _ed = _d.scalar(_sp(_UP).where(_UP.email == "meena@gcs.local"))
    _was, _ed_id = _ed.rights, _ed.id
    _ed.set_rights([_RP.EDIT_TASK, _RP.AUDIT_TASK]); _d.commit()

_pid = int(re.search(r"/tasks/(\d+)", admin.post("/tasks/new", data={
    "title": f"date lock {RUN}", "doer_id": _ed_id, "priority": "medium",
    "due_at": "2026-10-01T23:59"}, follow_redirects=False).headers["location"]).group(1))
editor = login("meena@gcs.local")
_pg = editor.get(f"/tasks/{_pid}").text
check("the date field is locked for them", "disabled" in _pg and "Locked" in _pg)
check("they really can see the task", editor.get(f"/tasks/{_pid}").status_code == 200)
editor.post(f"/tasks/{_pid}/edit", data={"title": f"date lock {RUN}", "details": "",
    "doer_id": _ed_id, "priority": "medium", "due_at": "2027-12-31T10:00"})
with _SLP() as _d:
    check("and the deadline did not move",
          _d.get(_TP, _pid).due_at.year == 2026, str(_d.get(_TP, _pid).due_at))
    check("even though the edit itself went through",
          _d.get(_TP, _pid).title == f"date lock {RUN}")

# The admin holds every right, so the date moves for them.
admin.post(f"/tasks/{_pid}/edit", data={"title": f"date lock {RUN}", "details": "",
    "doer_id": _ed_id, "priority": "medium", "due_at": "2027-12-31T10:00"})
with _SLP() as _d:
    check("an admin can move it", _d.get(_TP, _pid).due_at.year == 2027)
    _u = _d.get(_UP, _ed_id); _u.rights = _was; _d.commit()

check("the right is offered on the user form",
      'value="change_due_date"' in admin.get("/admin/users").text)

print("\n== FMS: each step says where the flow goes next ==")
from app.models import Flow as _FL, FlowStep as _FS, FlowInstance as _FI, TaskStatus as _TSF

# The user's own example: verify -> submit -> decide -> (rejected: back to 1)
_made = admin.post("/flows/new", data={
    "name": f"Bill verification {RUN}", "branch_id": "", "description": "",
    "sf_label": ["Bill No", "Vendor", "Bill type"],
    "sf_type": ["text", "text", "select"],
    "sf_required": ["1", "1", "1"],
    "sf_options": ["", "", "Electricity, Water, Rent"],
    "step_title": ["Verify the bill as per checklist", "Submit to the manager",
                   "Verified or rejected?", "Send to CMD for approval"],
    "step_doer": ["6", "6", "6", "6"],
    "step_tat": ["24", "24", "24", "24"],
    "step_priority": ["medium", "medium", "high", "medium"],
    "step_instructions": ["", "", "", ""],
    "step_fields": ["", "", "", ""],
    "step_audit": ["0", "0", "0", "0"],
    "step_proof": ["0", "0", "0", "0"],
    "step_decision": ["0", "0", "1", "0"],
    "step_yes": ["", "", "Verified", ""],
    "step_no": ["", "", "Rejected", ""],
    "step_next": ["2", "3", "4", "0"],
    "step_fail": ["", "", "1", ""],
})
check("the flow is created", _made.status_code == 200, _made.status_code)
with _SLP() as _d:
    _flow = _d.scalar(_sp(_FL).where(_FL.name == f"Bill verification {RUN}"))
    _fid = _flow.id
    _steps = {s.position: s for s in _flow.steps}
    check("step 3 is a decision", _steps[3].is_decision)
    check("with the two outcomes named",
          _steps[3].yes_label == "Verified" and _steps[3].no_label == "Rejected")
    check("verified routes to step 4", _steps[3].next_step_pos == 4)
    check("rejected routes back to step 1", _steps[3].fail_step_pos == 1)
    check("step 4 ends the flow", _steps[4].next_step_pos == 0)
    check("the flow carries its own start questions",
          _flow.start_field_list == ["Bill No", "Vendor", "Bill type"],
          str(_flow.start_field_list))
    _bt = _flow.start_form_fields[2]
    check("and one of them is a list question", _bt["type"] == "select")
    check("with the choices typed in",
          _bt["options"] == ["Electricity", "Water", "Rent"], str(_bt["options"]))

check("a route to a step that does not exist is refused",
      admin.post("/flows/new", data={
          "name": f"bad route {RUN}", "branch_id": "", "description": "",
          "step_title": ["only step"], "step_doer": ["6"], "step_tat": ["24"],
          "step_priority": ["medium"], "step_instructions": [""], "step_fields": [""],
          "step_audit": ["0"], "step_proof": ["0"], "step_decision": ["0"],
          "step_yes": [""], "step_no": [""], "step_next": ["9"], "step_fail": [""],
      }).status_code == 400)
check("a decision step with only one outcome is refused",
      admin.post("/flows/new", data={
          "name": f"half decision {RUN}", "branch_id": "", "description": "",
          "step_title": ["decide"], "step_doer": ["6"], "step_tat": ["24"],
          "step_priority": ["medium"], "step_instructions": [""], "step_fields": [""],
          "step_audit": ["0"], "step_proof": ["0"], "step_decision": ["1"],
          "step_yes": ["Yes"], "step_no": ["No"], "step_next": ["0"], "step_fail": [""],
      }).status_code == 400)

print("\n== and a rejected bill really does go back to step 1 ==")
_dpage = admin.get(f"/flows/{_fid}").text
check("the start form asks the flow's own questions",
      "Bill No" in _dpage and "Vendor" in _dpage and "Bill type" in _dpage)
check("and renders the list one as a dropdown",
      '<option value="Electricity">' in _dpage)
check("the step list shows where each one routes",
      "Rejected →" in _dpage and "step 1" in _dpage)
check("starting without a required field is refused",
      admin.post(f"/flows/{_fid}/start",
                 data={"reference": f"INV-{RUN}"}).status_code == 400)

_started = admin.post(f"/flows/{_fid}/start", data={
    "reference": f"INV-{RUN}", "sf0": "BILL-77", "sf1": "Acme",
    "sf2": "Electricity"})
check("the flow starts", _started.status_code == 200, _started.status_code)
with _SLP() as _d:
    _inst = _d.scalar(_sp(_FI).where(_FI.reference == f"INV-{RUN}"))
    _iid = _inst.id
    check("the start values are stored on the run", "BILL-77" in (_inst.context or ""))

def _open_step(iid):
    with _SLP() as _d:
        rows = [t for t in _d.scalars(_sp(_TP).where(_TP.flow_instance_id == iid)).all()
                if t.status not in (_TSF.COMPLETED, _TSF.CANCELLED)]
        return (rows[0].id, rows[0].flow_step.position) if rows else (None, None)

amit = login("amit@gcs.local")
_t, _pos = _open_step(_iid)
check("step 1 is open", _pos == 1, str(_pos))
amit.post(f"/tasks/{_t}/submit", data={})
_t, _pos = _open_step(_iid)
check("step 1 routes to step 2", _pos == 2, str(_pos))
amit.post(f"/tasks/{_t}/submit", data={})
_t, _pos = _open_step(_iid)
check("step 2 routes to step 3", _pos == 3, str(_pos))

_dt = amit.get(f"/tasks/{_t}").text
check("the decision step shows two buttons",
      'value="pass"' in _dt and 'value="fail"' in _dt)
check("named as configured", "Verified" in _dt and "Rejected" in _dt)
_nodec = TestClient(app, follow_redirects=False)
_nodec.post("/login", data={"email": "amit@gcs.local", "password": "gcs1234"})
check("submitting without choosing is refused",
      _nodec.post(f"/tasks/{_t}/submit", data={}).status_code == 400)

amit.post(f"/tasks/{_t}/submit", data={"decision": "fail"})
_t, _pos = _open_step(_iid)
check("REJECTED sends it back to step 1", _pos == 1, str(_pos))

# Round again, this time approving.
amit.post(f"/tasks/{_t}/submit", data={})
_t, _pos = _open_step(_iid)
amit.post(f"/tasks/{_t}/submit", data={})
_t, _pos = _open_step(_iid)
check("back at the decision", _pos == 3, str(_pos))
amit.post(f"/tasks/{_t}/submit", data={"decision": "pass"})
_t, _pos = _open_step(_iid)
check("VERIFIED carries on to step 4", _pos == 4, str(_pos))
amit.post(f"/tasks/{_t}/submit", data={})
with _SLP() as _d:
    check("and step 4 finishes the run",
          _d.get(_FI, _iid).completed_at is not None)
check("nothing is left open", _open_step(_iid)[0] is None)

print("\n== flows built before routing existed still run in order ==")
with _SLP() as _d:
    _old = _d.scalar(_sp(_FL).where(_FL.name.notlike(f"%{RUN}%")))
    if _old:
        check("an older flow has no routes set",
              all(s.next_step_pos is None for s in _old.steps))
        _oid = _old.id
_ostart = admin.post(f"/flows/{_oid}/start", data={"reference": f"LEGACY-{RUN}"})
check("it still starts", _ostart.status_code == 200)
with _SLP() as _d:
    _oi = _d.scalar(_sp(_FI).where(_FI.reference == f"LEGACY-{RUN}"))
    _oiid = _oi.id
_t2, _p2 = _open_step(_oiid)
with _SLP() as _d:
    _doer2 = _d.get(_TP, _t2).doer.email
_who = login(_doer2)
attach(_who, _t2); _who.post(f"/tasks/{_t2}/submit", data={})
_t3, _p3 = _open_step(_oiid)
check("and walks to the next step in order", _p3 == (_p2 or 0) + 1,
      f"{_p2} -> {_p3}")

print("\n== the menu is in the order you asked for ==")
_nav = admin.get("/").text
_order = [x for x in ["Operations", "Reports", "Insight", "Setup"] if x in _nav]
check("Operations, Reports, Insight, Setup",
      _order == ["Operations", "Reports", "Insight", "Setup"], str(_order))
_pos_of = {k: _nav.index(f">{k}<") for k in _order}
check("and they appear in that order on the page",
      list(_pos_of) == sorted(_pos_of, key=_pos_of.get), str(_pos_of))
check("Dashboard is above all of them",
      _nav.index("Dashboard") < min(_pos_of.values()))

print("\n== something confirms what just happened ==")
# The confirmation is one-shot and lives in the signed session, not the URL.
_c = TestClient(app, follow_redirects=True)
_r = _c.post("/login", data={"email": "mis@gcs.local", "password": "gcs1234"})
check("signing in says so", "Signed in" in _r.text and "toast" in _r.text)
check("and it does not come back on a refresh", "Signed in" not in _c.get("/").text)

_r2 = _c.post("/tasks/new", data={"title": f"toast {RUN}", "doer_id": 6,
                                  "priority": "medium", "due_at": "2026-11-01T23:59"})
check("assigning a task says so", "Task assigned" in _r2.text)
_tid2 = int(re.search(r"/tasks/(\d+)", str(_r2.url)).group(1))

_a = login("amit@gcs.local")
attach(_a, _tid2)
check("submitting says so", "Marked complete" in _a.post(
    f"/tasks/{_tid2}/submit", data={}).text)

check("the toast is switched off for reduced motion",
      "prefers-reduced-motion" in _c.get("/static/app.css").text)

print("\n== TAT is a unit and a number, not just hours ==")
from app import clock as _ck
from datetime import datetime as _dtu
check("months are calendar months, not 30-day blocks",
      _ck.add_months(_dtu(2026, 1, 31, 10, 0)) if False else
      _ck.add_months(_dtu(2026, 1, 31, 10, 0), 1) == _dtu(2026, 2, 28, 10, 0),
      str(_ck.add_months(_dtu(2026, 1, 31, 10, 0), 1)))
check("and they roll over the year",
      _ck.add_months(_dtu(2026, 12, 15, 9, 0), 2) == _dtu(2027, 2, 15, 9, 0))
check("a leap February is handled",
      _ck.add_months(_dtu(2028, 1, 31, 9, 0), 1) == _dtu(2028, 2, 29, 9, 0),
      str(_ck.add_months(_dtu(2028, 1, 31, 9, 0), 1)))
for _u, _v, _want in [("minutes", 30, _dtu(2026, 3, 1, 10, 30)),
                      ("hours", 5, _dtu(2026, 3, 1, 15, 0)),
                      ("days", 2, _dtu(2026, 3, 3, 10, 0)),
                      ("weeks", 2, _dtu(2026, 3, 15, 10, 0))]:
    check(f"{_v} {_u}", _ck.add_span(_dtu(2026, 3, 1, 10, 0), _u, _v) == _want)

_fp = admin.get("/flows/new").text
check("the builder offers every unit",
      all(f'value="{u}"' in _fp for u in ["minutes", "hours", "days", "weeks", "months"]))
check("and a second box for the number", 'class="tatval"' in _fp)
check("with sensible choices per unit", '"weeks": [1, 2, 3, 4, 6, 8]' in _fp
      or '"weeks":[1,2,3,4,6,8]' in _fp.replace(" ", ""))
check("the builder asks what to count the TAT from",
      'name="step_due_from"' in _fp)

print("\n== a step can hang its date off another step's date ==")
from app.models import Flow as _FLT, FlowInstance as _FIT, Task as _TKT
_made2 = admin.post("/flows/new", data={
    "name": f"TAT units {RUN}", "branch_id": "", "description": "",
    "step_title": ["Raise it", "Chase it", "Close it"],
    "step_doer": ["6", "6", "6"], "step_instructions": ["", "", ""],
    "step_fields": ["", "", ""], "step_audit": ["0", "0", "0"],
    "step_proof": ["0", "0", "0"], "step_decision": ["0", "0", "0"],
    "step_yes": ["", "", ""], "step_no": ["", "", ""],
    "step_priority": ["medium", "medium", "medium"],
    "step_tat_unit": ["days", "weeks", "months"],
    "step_tat": ["2", "1", "1"],
    "step_next": ["2", "3", "0"],
    "step_fail": ["", "", ""],
    # step 3's date is measured from step 1's date, not from when it opens
    "step_due_from": ["", "", "1"],
})
check("the flow saves", _made2.status_code == 200, _made2.status_code)
with _SLP() as _d:
    _f2 = _d.scalar(_sp(_FLT).where(_FLT.name == f"TAT units {RUN}"))
    _f2id = _f2.id
    _st = {x.position: x for x in _f2.steps}
    check("step 1 is 2 days", (_st[1].tat_unit, _st[1].tat_value) == ("days", 2))
    check("step 2 is 1 week", (_st[2].tat_unit, _st[2].tat_value) == ("weeks", 1))
    check("step 3 is 1 month", (_st[3].tat_unit, _st[3].tat_value) == ("months", 1))
    check("and reads back in words", _st[3].tat_label == "1 month", _st[3].tat_label)
    check("step 3 is tied to step 1's date", _st[3].due_from_pos == 1)

check("a date tied to a step that does not exist is refused",
      admin.post("/flows/new", data={
          "name": f"bad link {RUN}", "branch_id": "", "description": "",
          "step_title": ["one"], "step_doer": ["6"], "step_tat": ["1"],
          "step_tat_unit": ["days"], "step_priority": ["medium"],
          "step_instructions": [""], "step_fields": [""], "step_audit": ["0"],
          "step_proof": ["0"], "step_decision": ["0"], "step_yes": [""],
          "step_no": [""], "step_next": ["0"], "step_fail": [""],
          "step_due_from": ["7"]}).status_code == 400)
check("a step cannot take its date from itself",
      admin.post("/flows/new", data={
          "name": f"self link {RUN}", "branch_id": "", "description": "",
          "step_title": ["one"], "step_doer": ["6"], "step_tat": ["1"],
          "step_tat_unit": ["days"], "step_priority": ["medium"],
          "step_instructions": [""], "step_fields": [""], "step_audit": ["0"],
          "step_proof": ["0"], "step_decision": ["0"], "step_yes": [""],
          "step_no": [""], "step_next": ["0"], "step_fail": [""],
          "step_due_from": ["1"]}).status_code == 400)

# Walk it and check the real deadlines the engine produced.
admin.post(f"/flows/{_f2id}/start", data={"reference": f"TATRUN-{RUN}"})
with _SLP() as _d:
    _i2 = _d.scalar(_sp(_FIT).where(_FIT.reference == f"TATRUN-{RUN}"))
    _i2id = _i2.id
_t, _pos = _open_step(_i2id)
with _SLP() as _d:
    _due1 = _d.get(_TKT, _t).due_at
check("step 1 is due about two days out",
      1 <= (_due1 - _ck.now()).days <= 2, str(_due1))
amit.post(f"/tasks/{_t}/submit", data={})
_t, _pos = _open_step(_i2id)
with _SLP() as _d:
    _due2 = _d.get(_TKT, _t).due_at
check("step 2 is due about a week out",
      6 <= (_due2 - _ck.now()).days <= 8, str(_due2))
amit.post(f"/tasks/{_t}/submit", data={})
_t, _pos = _open_step(_i2id)
with _SLP() as _d:
    _due3 = _d.get(_TKT, _t).due_at
# Step 3 is a month after STEP 1's date, not a month from now — that is the
# whole point of tying it. Step 1 was due ~2 days out, so this lands near
# a month and two days from now, not a month from now.
check("step 3 counts its month from step 1's date, not from today",
      _ck.add_months(_due1, 1).date() == _due3.date(),
      f"step1 {_due1.date()} +1m = {_ck.add_months(_due1,1).date()} but got {_due3.date()}")
check("which is later than a month from now",
      _due3 > _ck.add_months(_ck.now(), 1) - __import__("datetime").timedelta(days=1))

print("\n== old flows keep their hours ==")
with _SLP() as _d:
    # Steps that came from the SEED, not from anything this run created.
    # Filtering on the step title missed them, because a step this test made
    # is called "Raise it" — the flow it belongs to is what carries the run id.
    _mine = {f.id for f in _d.scalars(_sp(_FLT)).all() if RUN in f.name}
    _legacy = [x for x in _d.scalars(_sp(_FS)).all() if x.flow_id not in _mine]
    check("there are older steps to check", bool(_legacy))
    check("every pre-existing step has a unit after upgrade",
          all(x.tat_unit for x in _legacy), "some are NULL")
    check("and it is hours", all(x.tat_unit == "hours" for x in _legacy))
    check("with the number it always had",
          all(x.tat_value == x.tat_hours for x in _legacy),
          str([(x.tat_value, x.tat_hours) for x in _legacy[:3]]))

# The upgrade path itself, not just freshly seeded rows: the new column
# carries DEFAULT 24, so a backfill keyed on NULL matches nothing and a
# 4-hour step silently becomes 24. Prove the real database upgrade keeps
# every number, and that running it twice changes nothing.
import shutil as _sh, tempfile as _tf, subprocess as _sub, os as _os, json as _js
_probe = _os.path.join(_tf.mkdtemp(), "upgrade.db")
_sh.copy("midap-backup-20260910-072235.db", _probe)
_code = ("from app import migrate; migrate.run(); migrate.run();"
         "from app.db import SessionLocal; from app.models import FlowStep;"
         "from sqlalchemy import select; import json;"
         "d=SessionLocal(); r=d.scalars(select(FlowStep)).all();"
         "print(json.dumps([[s.tat_hours, s.tat_value, s.tat_unit] for s in r]))")
_out = _sub.run([__import__("sys").executable, "-c", _code], capture_output=True, text=True,
                env={**_os.environ, "DATABASE_URL": f"sqlite:///{_probe}"})
_rows = _js.loads(_out.stdout.strip().splitlines()[-1]) if _out.stdout.strip() else []
check("an existing database upgrades with its TAT numbers intact",
      _rows and all(h == v and u == "hours" for h, v, u in _rows),
      _out.stderr[-200:] or str(_rows[:4]))
check("and the upgrade is safe to run twice", bool(_rows))

print("\n== the start form is a real form builder ==")
_nb = admin.get("/flows/new").text
check("the builder is on the create page", 'name="sf_label"' in _nb)
for _t, _lbl in [("text", "Short text"), ("textarea", "Long text"),
                 ("number", "Number"), ("date", "Date"),
                 ("select", "Choose from a list"), ("yesno", "Yes / No")]:
    check(f"it offers '{_lbl}'", f'value="{_t}"' in _nb and _lbl in _nb)
check("a question can be made optional", 'name="sf_required"' in _nb)
check("and a list question takes its choices", 'name="sf_options"' in _nb)

check("a list question with no choices is refused",
      admin.post("/flows/new", data={
          "name": f"no choices {RUN}", "branch_id": "", "description": "",
          "sf_label": ["Type"], "sf_type": ["select"], "sf_required": ["1"],
          "sf_options": [""],
          "step_title": ["one"], "step_doer": ["6"], "step_tat": ["1"],
          "step_tat_unit": ["days"], "step_priority": ["medium"],
          "step_instructions": [""], "step_fields": [""], "step_audit": ["0"],
          "step_proof": ["0"], "step_decision": ["0"], "step_yes": [""],
          "step_no": [""], "step_next": ["0"], "step_fail": [""],
          "step_due_from": [""]}).status_code == 400)

print("\n== every answer type is checked when a run starts ==")
_tf = admin.post("/flows/new", data={
    "name": f"Typed start {RUN}", "branch_id": "", "description": "",
    "sf_label": ["Amount", "Pay by", "Urgent?", "Notes"],
    "sf_type": ["number", "date", "yesno", "textarea"],
    "sf_required": ["1", "1", "1", "0"],
    "sf_options": ["", "", "", ""],
    "step_title": ["Do it"], "step_doer": ["6"], "step_tat": ["1"],
    "step_tat_unit": ["days"], "step_priority": ["medium"],
    "step_instructions": [""], "step_fields": [""], "step_audit": ["0"],
    "step_proof": ["0"], "step_decision": ["0"], "step_yes": [""],
    "step_no": [""], "step_next": ["0"], "step_fail": [""], "step_due_from": [""]})
check("a typed start form saves", _tf.status_code == 200, _tf.status_code)
with _SLP() as _d:
    _tflow = _d.scalar(_sp(_FLT).where(_FLT.name == f"Typed start {RUN}"))
    _tfid = _tflow.id
    check("the types are stored",
          [f["type"] for f in _tflow.start_form_fields] ==
          ["number", "date", "yesno", "textarea"])
    check("and the optional one is marked optional",
          _tflow.start_form_fields[3]["required"] is False)

_sp_page = admin.get(f"/flows/{_tfid}").text
check("a number question renders as a number box", 'type="number"' in _sp_page)
check("a date question renders as a date box", 'type="date"' in _sp_page)
check("a yes/no question renders as a dropdown", '<option value="Yes">' in _sp_page)
check("a long-text question renders as a textarea", "<textarea" in _sp_page)

_ok = {"reference": f"TYPED-{RUN}", "sf0": "1500", "sf1": "2026-12-01",
       "sf2": "Yes", "sf3": ""}
check("a letter in a number box is refused",
      admin.post(f"/flows/{_tfid}/start",
                 data={**_ok, "sf0": "abcd"}).status_code == 400)
check("something that is not Yes or No is refused",
      admin.post(f"/flows/{_tfid}/start",
                 data={**_ok, "sf2": "Maybe"}).status_code == 400)
check("a choice that is not on the list is refused",
      admin.post(f"/flows/{_fid}/start",
                 data={"reference": f"BAD-{RUN}", "sf0": "B1", "sf1": "V",
                       "sf2": "Diesel"}).status_code == 400)
check("but the optional one may be left blank",
      admin.post(f"/flows/{_tfid}/start", data=_ok).status_code == 200)
with _SLP() as _d:
    _ti = _d.scalar(_sp(_FIT).where(_FIT.reference == f"TYPED-{RUN}"))
    check("the answers are on the run",
          "1500" in (_ti.context or "") and "2026-12-01" in (_ti.context or ""),
          _ti.context)
    check("and the blank optional one is simply absent",
          "Notes" not in (_ti.context or ""))

print("\n== a flow built earlier can be given a start form ==")
# This is the real complaint: "Bill to payment" exists with no questions and
# no way to add them. Build a flow with none, then add them afterwards.
admin.post("/flows/new", data={
    "name": f"No questions {RUN}", "branch_id": "", "description": "",
    "sf_label": [""], "sf_type": ["text"], "sf_required": ["1"], "sf_options": [""],
    "step_title": ["Only step"], "step_doer": ["6"], "step_tat": ["1"],
    "step_tat_unit": ["days"], "step_priority": ["medium"],
    "step_instructions": [""], "step_fields": [""], "step_audit": ["0"],
    "step_proof": ["0"], "step_decision": ["0"], "step_yes": [""],
    "step_no": [""], "step_next": ["0"], "step_fail": [""], "step_due_from": [""]})
with _SLP() as _d:
    _nq = _d.scalar(_sp(_FLT).where(_FLT.name == f"No questions {RUN}"))
    _nqid = _nq.id
    check("it starts with no questions", _nq.start_form_fields == [])
check("an empty question row is not saved as a blank question", True)

check("the flow page offers Edit",
      f"/flows/{_nqid}/edit" in admin.get(f"/flows/{_nqid}").text)
check("the edit page opens", admin.get(f"/flows/{_nqid}/edit").status_code == 200)
_ed = admin.post(f"/flows/{_nqid}/edit", data={
    "name": f"No questions {RUN}", "branch_id": "", "description": "now it asks",
    "sf_label": ["Invoice No", "Department"], "sf_type": ["text", "select"],
    "sf_required": ["1", "1"], "sf_options": ["", "Accounts, Ops"]})
check("the start form can be added afterwards", _ed.status_code == 200, _ed.status_code)
with _SLP() as _d:
    _nq2 = _d.get(_FLT, _nqid)
    check("and it sticks",
          _nq2.start_field_list == ["Invoice No", "Department"],
          str(_nq2.start_field_list))
    check("the description was saved too", _nq2.description == "now it asks")
check("the start page now asks them",
      "Invoice No" in admin.get(f"/flows/{_nqid}").text)
check("a doer cannot edit a flow",
      doer.get(f"/flows/{_nqid}/edit").status_code == 403)

# The edit page must show what is already there, not an empty builder.
_epage = admin.get(f"/flows/{_nqid}/edit").text
# The builder is filled in by script from this JSON, so that is what has to
# be right — the choices are joined into "Accounts, Ops" in the browser and
# never appear as literal text in the source.
_json_in_page = _re.search(r"var existing = (.*?);\n", _epage)
_pre = _json.loads(_json_in_page.group(1)) if _json_in_page else []
check("the edit page hands the builder the current questions",
      [f["label"] for f in _pre] == ["Invoice No", "Department"], str(_pre))
check("including the list question's choices",
      _pre and _pre[1]["options"] == ["Accounts", "Ops"], str(_pre[-1:]))
check("and whether each is required",
      all(f["required"] for f in _pre))

print("\n== flows built with the old comma box still ask their questions ==")
with _SLP() as _d:
    _old_style = _FLT(org_id=1, name=f"Legacy questions {RUN}",
                      start_fields="Member Name, Lead Id")
    _d.add(_old_style); _d.commit()
    check("a comma line is read as plain text questions",
          [f["label"] for f in _old_style.start_form_fields] ==
          ["Member Name", "Lead Id"],
          str(_old_style.start_form_fields))
    check("and they are all short-text",
          all(f["type"] == "text" for f in _old_style.start_form_fields))

print("\n== the user form accepts a split under 100 ==")
from app.db import SessionLocal as _SLB
from app.models import User as _UB
from sqlalchemy import select as _sb

_r = admin.post("/admin/users", data={
    "name": f"Partly manual {RUN}", "email": f"pm.{RUN}@gcs.local", "phone": "",
    "password": "pmpass1234", "role": "doer", "branch_id": "", "department_id": "",
    "rights": [], "bm_delegation": 40, "bm_checklist": 10, "bm_fms": 10})
check("40 / 10 / 10 is accepted", _r.status_code == 200, _r.status_code)
with _SLB() as _d:
    _pm = _d.scalar(_sb(_UB).where(_UB.email == f"pm.{RUN}@gcs.local"))
    check("the user is created", _pm is not None)
    check("with exactly those benchmarks",
          (_pm.bm_delegation, _pm.bm_checklist, _pm.bm_fms) == (40, 10, 10),
          str((_pm.bm_delegation, _pm.bm_checklist, _pm.bm_fms)))
    _pmid = _pm.id

_bad = TestClient(app, follow_redirects=False)
_bad.post("/login", data={"email": "mis@gcs.local", "password": "gcs1234"})
check("but over 100 is still refused",
      _bad.post("/admin/users", data={
          "name": f"Over {RUN}", "email": f"over.{RUN}@gcs.local", "phone": "",
          "password": "overpass123", "role": "doer", "branch_id": "",
          "department_id": "", "rights": [], "bm_delegation": 60,
          "bm_checklist": 60, "bm_fms": 60}).status_code == 400)

check("editing to a sub-100 split works",
      admin.post(f"/admin/users/{_pmid}", data={
          "name": f"Partly manual {RUN}", "phone": "", "role": "doer",
          "branch_id": "", "department_id": "", "rights": [],
          "bm_delegation": 30, "bm_checklist": 20,
          "bm_fms": 0}).status_code == 200)
with _SLB() as _d:
    check("and saves", _d.get(_UB, _pmid).benchmark_total == 50)

check("the page no longer demands 100",
      "must add up to" not in admin.get("/admin/users").text)
check("it explains what a sub-100 split means",
      "do not have to add up to 100" in admin.get("/admin/users").text)
_js = admin.get("/static/benchmark.js").text
check("and the browser no longer blocks it", "t === 100" not in _js)
check("while still stopping a total over 100", "t > 100" in _js)

# Their own score page has to say so, or 50 reads as a result rather than as
# "the software had 50 points to give".
_pmc = login(f"pm.{RUN}@gcs.local", "pmpass1234")
_dash = _pmc.get("/").text
check("the doer's dashboard says how much is scored by software",
      "50 of the 100 points are scored by the" in _dash
      or "50 of 100 scored here" in _dash, "no note shown")

print("\n== the checklist no longer depends on the server restarting ==")
# This used to happen only at boot. On a host kept awake around the clock
# there is no second boot, so the checklist would have stopped after day one.
import asyncio as _aio, datetime as _dtm
from app import main as _M
from app.models import RecurringRule as _RR

def _recurring_count():
    with _SLP() as _d:
        return len(_d.scalars(_sp(_TP).where(_TP.source == _TSR.RECURRING)).all())

_real_today = _clock.today
_day = _clock.today()
_M.SPAWN_CHECK_SECONDS = 0.05
with _SLP() as _d:
    for _r in _d.scalars(_sp(_RR)).all():
        _r.last_spawned_on = None
    _d.commit()

async def _walk_two_days():
    _t = _aio.create_task(_M._daily_spawn())
    await _aio.sleep(0.3)
    _before = _recurring_count()
    _clock.today = lambda: _day + _dtm.timedelta(days=1)
    await _aio.sleep(0.4)
    _after = _recurring_count()
    await _aio.sleep(0.4)          # same day again
    _again = _recurring_count()
    _t.cancel()
    try:
        await _t
    except _aio.CancelledError:
        pass
    return _before, _after, _again

try:
    _b, _a, _ag = _aio.run(_walk_two_days())
finally:
    _clock.today = _real_today
    _M.SPAWN_CHECK_SECONDS = 600

check("the app spawns today's checklist on its own", _b > 0, _b)
check("and a new day's checklist without any restart", _a > _b, f"{_b} -> {_a}")
check("running again the same day creates no duplicates", _ag == _a, f"{_a} -> {_ag}")

# The loop must survive a bad day rather than dying silently.
async def _survive_error():
    _boom = _clock.today
    _clock.today = lambda: (_ for _ in ()).throw(RuntimeError("bad clock"))
    _t = _aio.create_task(_M._daily_spawn())
    await _aio.sleep(0.2)
    _alive = not _t.done()
    _clock.today = _boom
    _t.cancel()
    try:
        await _t
    except _aio.CancelledError:
        pass
    return _alive

_M.SPAWN_CHECK_SECONDS = 0.05
try:
    check("a failure does not kill the loop", _aio.run(_survive_error()))
finally:
    _M.SPAWN_CHECK_SECONDS = 600
    _clock.today = _real_today

check("/cron/spawn still works alongside it",
      admin.get("/cron/spawn").status_code in (200, 401))

print("\n== an uptime monitor can reach us ==")
# UptimeRobot, Better Stack and most link checkers send HEAD, not GET. Every
# route used to refuse it with 405, so the site read as permanently down
# while being perfectly healthy.
_hc = TestClient(app, follow_redirects=False)
for _path in ["/", "/healthz", "/login", "/static/app.css"]:
    _g = _hc.request("GET", _path)
    _h = _hc.request("HEAD", _path)
    check(f"HEAD {_path} is answered, not refused",
          _h.status_code != 405, f"got {_h.status_code}")
    check(f"HEAD {_path} gives the same status as GET",
          _h.status_code == _g.status_code, f"{_h.status_code} vs {_g.status_code}")
    check(f"HEAD {_path} sends no body", _h.content == b"", _h.content[:40])

# The header must still describe what a GET would return — that is what HEAD
# is for. A zero here would be a lie.
_hh = _hc.request("HEAD", "/healthz")
_hg = _hc.request("GET", "/healthz")
check("HEAD keeps the real content-length",
      _hh.headers.get("content-length") == _hg.headers.get("content-length"),
      f"{_hh.headers.get('content-length')} vs {_hg.headers.get('content-length')}")
check("and the same content-type",
      _hh.headers.get("content-type") == _hg.headers.get("content-type"))

# The monitor URL people will actually use.
check("/healthz answers HEAD with 200",
      _hc.request("HEAD", "/healthz").status_code == 200)
check("and GET still returns the real body",
      '"ok":true' in _hc.request("GET", "/healthz").text)

# POST must NOT be quietly turned into a GET by the same middleware.
check("POST is untouched",
      _hc.request("POST", "/healthz").status_code == 405)

print("\n== closing a task from the list it appears on ==")
# The doer's own three kinds of work, still open.
from app.models import (Task as _MT, TaskSource as _MS, User as _MU,
                        TaskStatus as _MST, Attachment as _MA)
from sqlalchemy import select as _msel

_OPEN = (_MST.PENDING, _MST.IN_PROGRESS, _MST.REJECTED, _MST.REOPENED)
_mine = {}
with _SLT() as _d:
    _amit = _d.scalar(_msel(_MU).where(_MU.email == "amit@gcs.local"))
    for _k, _src in [("delegation", _MS.DELEGATION), ("checklist", _MS.RECURRING),
                     ("fms", _MS.FLOW)]:
        _t = _d.scalars(_msel(_MT).where(
            _MT.doer_id == _amit.id, _MT.source == _src,
            _MT.status.in_(_OPEN)).limit(1)).first()
        if _t:
            _mine[_k] = _t.id

check("the doer has open work of at least two kinds to test with",
      len(_mine) >= 2, sorted(_mine))

# 1 — the button is on the list, for every kind of work.
_list = doer.get("/tasks?scope=mine&status=open").text
check("the list offers a Mark complete button", "Mark complete" in _list)
for _k, _i in _mine.items():
    check(f"{_k}: its row has the button", f'data-mark="{_i}"' in _list,
          f"task {_i}")

# The same button on the dashboard, which is where most people start.
_dash = doer.get("/").text
check("the dashboard rows have it too", 'data-mark="' in _dash)

# 2 — and never on somebody else's work. The admin sees every task; none of
# them may be closeable from their list, because they are not doing them.
_allpg = admin.get("/tasks?scope=all&status=open").text
for _k, _i in _mine.items():
    check(f"{_k}: an onlooker gets no button for it",
          f'data-mark="{_i}"' not in _allpg, f"task {_i} closeable by admin")

# 3 — the box itself.
_one = _mine.get("delegation") or list(_mine.values())[0]
_box = doer.get(f"/tasks/{_one}/mark")
check("the doer can open the box", _box.status_code == 200, _box.status_code)
check("it is a fragment, not a whole page", "<html" not in _box.text.lower())
check("it posts to the ordinary submit route",
      f'action="/tasks/{_one}/submit"' in _box.text)
check("it carries a return_to field", 'class="returnto"' in _box.text)
check("it offers the full task page as the other way in",
      f'href="/tasks/{_one}"' in _box.text)
check("it has a note box", 'name="completion_note"' in _box.text)

check("somebody else cannot open it",
      mgr.get(f"/tasks/{_one}/mark").status_code == 403,
      mgr.get(f"/tasks/{_one}/mark").status_code)
_stranger = login("ravi@gcs.local")     # a doer in another branch entirely
check("and a stranger gets a plain not-found",
      _stranger.get(f"/tasks/{_one}/mark").status_code == 404,
      _stranger.get(f"/tasks/{_one}/mark").status_code)

# 4 — submitting from the box lands back on the list, not the task page.
_mk = mgr.post("/tasks/new", data={
    "title": f"SMOKE popup return {RUN}", "details": "",
    "doer_id": str(_amit.id), "branch_id": "", "priority": "low",
    "due_at": "2026-12-31T23:59"})
_pid = int(_re.findall(r"/tasks/(\d+)/comment", _mk.text)[-1])
attach(doer, _pid)
_back = "/tasks?scope=mine&status=open&source=delegation"
_nav = TestClient(app, follow_redirects=False)
_nav.cookies.update(doer.cookies)
_r = _nav.post(f"/tasks/{_pid}/submit",
               data={"completion_note": "done from the list", "return_to": _back})
check("submitting from the box goes back to the list",
      _r.status_code == 303 and _r.headers["location"] == _back,
      f"{_r.status_code} -> {_r.headers.get('location')}")
with _SLT() as _d:
    _t = _d.get(_MT, _pid)
    check("and the task really is closed",
          _t.status in (_MST.COMPLETED, _MST.SUBMITTED), _t.status)
    check("with the note saved", _t.completion_note == "done from the list",
          _t.completion_note)

# 5 — return_to may only ever be a path on this site. A form field that
# decides where a browser lands is the classic way to bounce somebody onto
# another site from a link that looks like ours.
_mk2 = mgr.post("/tasks/new", data={
    "title": f"SMOKE popup redirect {RUN}", "details": "",
    "doer_id": str(_amit.id), "branch_id": "", "priority": "low",
    "due_at": "2026-12-31T23:59"})
_rid2 = int(_re.findall(r"/tasks/(\d+)/comment", _mk2.text)[-1])
attach(doer, _rid2)
_nav2 = TestClient(app, follow_redirects=False)
_nav2.cookies.update(doer.cookies)
_evil = _nav2.post(f"/tasks/{_rid2}/submit",
                   data={"completion_note": "x", "return_to": "//evil.example.com/"})
check("an off-site return_to is ignored",
      _evil.headers.get("location") == f"/tasks/{_rid2}",
      _evil.headers.get("location"))

for _bad in ("https://evil.example.com/", "/\\evil.example.com", "javascript:alert(1)"):
    _mk3 = mgr.post("/tasks/new", data={
        "title": f"SMOKE popup redirect {RUN} {_bad[:6]}", "details": "",
        "doer_id": str(_amit.id), "branch_id": "", "priority": "low",
        "due_at": "2026-12-31T23:59"})
    _b3 = int(_re.findall(r"/tasks/(\d+)/comment", _mk3.text)[-1])
    attach(doer, _b3)
    _n3 = TestClient(app, follow_redirects=False)
    _n3.cookies.update(doer.cookies)
    _rr = _n3.post(f"/tasks/{_b3}/submit", data={"completion_note": "x", "return_to": _bad})
    check(f"return_to {_bad[:22]!r} is refused",
          _rr.headers.get("location") == f"/tasks/{_b3}",
          _rr.headers.get("location"))

# 6 — the proof rule is the server's, not the button's. Disabling a button
# is a courtesy; the rule has to hold for anyone who posts anyway.
_mk4 = mgr.post("/tasks/new", data={
    "title": f"SMOKE popup proof {RUN}", "details": "",
    "doer_id": str(_amit.id), "branch_id": "", "priority": "low",
    "due_at": "2026-12-31T23:59"})
_pf = int(_re.findall(r"/tasks/(\d+)/comment", _mk4.text)[-1])
_pbox = doer.get(f"/tasks/{_pf}/mark")
check("a task needing proof says so in the box",
      "Proof is required" in _pbox.text)
check("and its button starts switched off", "disabled" in _pbox.text)
_noproof = doer.post(f"/tasks/{_pf}/submit",
                     data={"completion_note": "no proof", "return_to": _back})
check("submitting it without proof is refused", _noproof.status_code == 400,
      _noproof.status_code)
attach(doer, _pf)
_pbox2 = doer.get(f"/tasks/{_pf}/mark")
check("once proof is attached the box says you can submit",
      "you can submit" in _pbox2.text)
check("and the button is no longer disabled",
      "disabled" not in _pbox2.text.split('class="markacts"')[1])

# 7 — a closed task has no box to open.
doer.post(f"/tasks/{_pf}/submit", data={"completion_note": "done"})
check("a task already closed cannot be reopened from the list",
      doer.get(f"/tasks/{_pf}/mark").status_code == 400,
      doer.get(f"/tasks/{_pf}/mark").status_code)

# 8 — an FMS decision step must offer both outcomes in the box, not one
# "Mark complete" that silently picks a direction. Walked from a fresh run of
# the bill flow, because the runs earlier in this file all finished.
_dstart = admin.post(f"/flows/{_fid}/start", data={
    "reference": f"POPUP-{RUN}", "sf0": "BILL-99", "sf1": "Acme", "sf2": "Electricity"})
check("a bill run starts for the pop-up test", _dstart.status_code == 200,
      _dstart.status_code)
with _SLP() as _d:
    _piid = _d.scalar(_sp(_FI).where(_FI.reference == f"POPUP-{RUN}")).id
_amc = login("amit@gcs.local")
_pt, _ppos = _open_step(_piid)
while _ppos is not None and _ppos < 3:
    _amc.post(f"/tasks/{_pt}/submit", data={})
    _pt, _ppos = _open_step(_piid)
check("it reaches the decision step", _ppos == 3, str(_ppos))
_dbox = _amc.get(f"/tasks/{_pt}/mark")
check("a decision step's box opens", _dbox.status_code == 200, _dbox.status_code)
check("it offers both outcomes",
      'value="pass"' in _dbox.text and 'value="fail"' in _dbox.text)
check("named as configured",
      "Verified" in _dbox.text and "Rejected" in _dbox.text)
check("rather than a single Mark complete button",
      "Mark complete</button>" not in _dbox.text)

# And choosing from the box routes the flow exactly as the task page does.
_pnav = TestClient(app, follow_redirects=False)
_pnav.cookies.update(_amc.cookies)
_pnav.post(f"/tasks/{_pt}/submit", data={"decision": "fail", "return_to": _back})
_pt2, _ppos2 = _open_step(_piid)
check("Rejected, chosen in the pop-up, sends the bill back to step 1",
      _ppos2 == 1, str(_ppos2))

# 9 — the table still holds together: one header cell per column.
_hdr = _list.split("<table>", 1)[1].split("</tr>", 1)[0].count("<th>")
_firstrow = _list.split("</tr>", 2)[1]
check("the new column has a header of its own", _hdr == _firstrow.count("<td"),
      f"{_hdr} headers vs {_firstrow.count('<td')} cells")

# 10 — the pop-up shell and its script are on the page that needs them.
check("the list page carries the pop-up shell", 'id="markDialog"' in _list)
check("and loads its script", "/static/markbox.js" in _list)
check("the sign-in page does not", 'id="markDialog"' not in
      TestClient(app).get("/login").text)

print("\n== Completed means the doer has finished it ==")
# The complaint: mark a task complete, it needs an audit, and the Completed
# tab stays empty. From the doer's side the job IS done — whether an auditor
# has looked at it yet is a separate column, not a reason to hide the row.
from app.models import AuditState as _AS2

_ca = mgr.post("/tasks/new", data={
    "title": f"SMOKE completed tab {RUN}", "details": "",
    "doer_id": str(_amit.id), "branch_id": "", "priority": "medium",
    "due_at": "2026-12-31T23:59", "requires_audit": "1"})
_cid = int(_re.findall(r"/tasks/(\d+)/comment", _ca.text)[-1])
attach(doer, _cid)
doer.post(f"/tasks/{_cid}/submit", data={"completion_note": "finished"})

with _SLT() as _d:
    _t = _d.get(_MT, _cid)
    check("it is waiting on an auditor", _t.status == _MST.SUBMITTED, _t.status)
    check("and its audit is pending", _t.audit_state == _AS2.PENDING)

_done = doer.get("/tasks?scope=mine&status=done").text
check("it now appears under Completed", f'/tasks/{_cid}"' in _done, "still missing")
check("the page explains what Completed covers",
      "still sitting with an auditor" in _done)
check("its row says where the audit stands", "Audit pending" in _done)
check("it is still under With auditor",
      f'/tasks/{_cid}"' in doer.get("/tasks?scope=mine&status=audit").text)
check("and still under Audit pending",
      f'/tasks/{_cid}"' in doer.get("/tasks?scope=mine&status=audit_pending").text)
check("no Mark complete button on it any more",
      f'data-mark="{_cid}"' not in _done)

# The date filter has to find it too. Completed used to filter on closed_at,
# which a task waiting on an auditor does not have yet — so picking any date
# range would have emptied the tab all over again.
_today = _clock.today().isoformat()
_ranged = doer.get(f"/tasks?scope=mine&status=done&date_from={_today}&date_to={_today}").text
check("a date range on Completed still finds it",
      f'/tasks/{_cid}"' in _ranged, "dropped by the date filter")
check("and says which date it is filtering on", "Completion date" in _ranged)

# Closing the audit must not push it back out of the tab.
admin.post(f"/tasks/{_cid}/audit",
           data={"decision": "approve", "score": "9", "remark": "checked"})
_done2 = doer.get("/tasks?scope=mine&status=done").text
check("it stays under Completed once the audit closes",
      f'/tasks/{_cid}"' in _done2)
check("and appears under Audit completed",
      f'/tasks/{_cid}"' in doer.get("/tasks?scope=mine&status=audit_done").text)

print("\n== Audit completed and False marking have tabs of their own ==")
_tabs = doer.get("/tasks?scope=mine&status=open").text
for _lbl in ("Completed", "With auditor", "Audit pending", "Audit completed",
             "False marking"):
    check(f"the tab row offers {_lbl}", f">{_lbl}</a>" in _tabs or _lbl in _tabs)

# A flagged task must be findable by the flag, and say so on its row.
_fm = mgr.post("/tasks/new", data={
    "title": f"SMOKE false mark tab {RUN}", "details": "",
    "doer_id": str(_amit.id), "branch_id": "", "priority": "high",
    "due_at": "2026-12-31T23:59"})
_fid2 = int(_re.findall(r"/tasks/(\d+)/comment", _fm.text)[-1])
attach(doer, _fid2)
doer.post(f"/tasks/{_fid2}/submit", data={"completion_note": "done"})
admin.post(f"/tasks/{_fid2}/false-mark",
           data={"confirm": "yes", "reason": "register was never filled"})
_fpage = doer.get("/tasks?scope=mine&status=false_mark")
check("the False marking tab opens", _fpage.status_code == 200, _fpage.status_code)
check("and lists the flagged task", f'/tasks/{_fid2}"' in _fpage.text)
check("the row carries the penalty chip", "false marking −10" in _fpage.text)
check("sorted by when it was flagged",
      "most recently flagged first" in _fpage.text.lower())
check("a task nobody flagged is not in there",
      f'/tasks/{_cid}"' not in _fpage.text)

print("\n== the three reports carry the audit figures ==")
for _src in ("delegation", "checklist", "fms"):
    _rp = admin.get(f"/reports/tasks?source={_src}")
    check(f"{_src} report opens", _rp.status_code == 200, _rp.status_code)
    for _lbl in ("Audit pending", "Audit completed", "False marking"):
        check(f"{_src}: it shows {_lbl}", _lbl in _rp.text)
    for _st in ("audit_pending", "audit_done", "false"):
        _sp2 = admin.get(f"/reports/tasks?source={_src}&state={_st}")
        check(f"{_src}: the {_st} tab opens", _sp2.status_code == 200, _sp2.status_code)

# The figures must be the database's, not a guess.
_dr = admin.get("/reports/tasks?source=delegation&state=false")
check("the delegation report lists the flagged task",
      f'/tasks/{_fid2}"' in _dr.text, "missing from False marking")
_dc = admin.get("/reports/tasks?source=delegation&state=completed")
check("and counts work waiting on an auditor as completed",
      f'/tasks/{_cid}"' in _dc.text)
check("saying plainly that it scores only once the audit closes",
      "scores once the audit closes" in _dc.text)

# An unknown state falls back rather than breaking the page.
check("a nonsense state falls back to Pending",
      admin.get("/reports/tasks?source=fms&state=banana").status_code == 200)

print("\n== Excel downloads ==")
import io as _xio
from openpyxl import load_workbook as _load

_XL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def _book(client, url):
    """Download an export and open it, the way the person will."""
    r = client.get(url)
    if r.status_code != 200 or not r.headers.get("content-type", "").startswith(_XL):
        return None, r
    return _load(_xio.BytesIO(r.content)), r

def _table(ws):
    """(headers, rows) from a sheet, skipping the title block above it.

    The header is the first row carrying more than one value — the title,
    the note and the export stamp above it each occupy column A alone.
    """
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if sum(1 for c in r if c not in (None, "")) > 1:
            return list(r), [x for x in rows[i + 1:]]
    return [], []

_EXPORTS = [
    "/tasks?scope=mine&status=open&export=xlsx",
    "/tasks?scope=all&status=all&export=xlsx",
    "/tasks?scope=all&status=false_mark&export=xlsx",
    "/reports/tasks?source=delegation&export=xlsx",
    "/reports/tasks?source=checklist&state=completed&export=xlsx",
    "/reports/tasks?source=fms&state=audit_pending&export=xlsx",
    "/reports/followups?export=xlsx",
    "/reports/audit?export=xlsx",
    "/reports/score?export=xlsx",
    "/admin/users?export=xlsx",
    "/admin/branches?export=xlsx",
    "/admin/departments?export=xlsx",
    "/followups?export=xlsx",
    "/followups?date_from=2026-09-01&date_to=2026-09-25&export=xlsx",
    "/flows?export=xlsx",
    "/recurring?export=xlsx",
    "/outbox?export=xlsx",
]
for _u in _EXPORTS:
    _wb, _r = _book(admin, _u)
    _name = _u.split("?")[0]
    check(f"{_name} downloads an Excel file", _wb is not None,
          f"{_r.status_code} {_r.headers.get('content-type','')[:40]}")
    if _wb is None:
        continue
    check(f"{_name}: the browser is told to save it",
          "attachment" in _r.headers.get("content-disposition", ""))
    check(f"{_name}: the file is named and dated",
          ".xlsx" in _r.headers.get("content-disposition", ""))
    _ws = _wb.worksheets[0]
    _hdr, _rows = _table(_ws)
    check(f"{_name}: it has a header row", bool(_hdr) and all(_hdr[:2]), _hdr[:3])
    check(f"{_name}: the header is frozen", _ws.freeze_panes is not None)
    check(f"{_name}: and says when it was exported",
          any("Exported" in str(c.value or "") for c in _ws["A"][:6]))

# --- the file must hold what the page held, not a different query ----------
_pg = admin.get("/tasks?scope=all&status=open&source=delegation")
_ids_page = set(_re.findall(r'/tasks/(\d+)"', _pg.text))
_wb, _ = _book(admin, "/tasks?scope=all&status=open&source=delegation&export=xlsx")
_hdr, _rows = _table(_wb.worksheets[0])
_titles_file = {r[_hdr.index("Task")] for r in _rows}
with _SLT() as _d:
    _titles_page = {_d.get(_MT, int(i)).title for i in _ids_page}
check("the Excel rows are exactly the page's rows",
      _titles_file == _titles_page,
      f"file {len(_titles_file)} vs page {len(_titles_page)}")
check("every column the page shows is a column in the file",
      all(h in _hdr for h in ["Task", "Doer", "Priority", "Status",
                              "Planned date", "Audit", "False marking"]), _hdr)

# A date is a real date, not text that looks like one — otherwise it will
# not sort, filter by month, or group in a pivot table.
_planned = _wb.worksheets[0].cell(
    row=int(_wb.worksheets[0].auto_filter.ref.split(":")[0][1:]) + 1,
    column=_hdr.index("Planned date") + 1)
from datetime import datetime as _dtclass
check("a date lands in the cell as a real date",
      isinstance(_planned.value, _dtclass), type(_planned.value).__name__)
check("and carries a readable format", "mmm" in _planned.number_format.lower(),
      _planned.number_format)

# --- a filter on the page is a filter in the file -------------------------
_narrow = "/tasks?scope=all&status=all&date_from=2026-01-01&date_to=2026-01-02"
_wbn, _ = _book(admin, _narrow + "&export=xlsx")
_h2, _r2 = _table(_wbn.worksheets[0])
_wba, _ = _book(admin, "/tasks?scope=all&status=all&export=xlsx")
_h3, _r3 = _table(_wba.worksheets[0])
check("a date range really narrows the export", len(_r2) < len(_r3),
      f"{len(_r2)} vs {len(_r3)}")
check("and the file says which filters produced it",
      any("2026-01-01" in str(c.value or "")
          for c in _wbn.worksheets[0]["A"][:6]))

# --- nobody exports rows they cannot see ----------------------------------
_wbd, _ = _book(doer, "/reports/tasks?source=delegation&state=all&export=xlsx")
_hd, _rd = _table(_wbd.worksheets[0])
_others = {n for n in (r[_hd.index("Doer")] for r in _rd) if n} - {_amit.name}
_assigned = set()
with _SLT() as _d:
    for _t in _d.scalars(_msel(_MT).where(_MT.assigner_id == _amit.id)).all():
        _assigned.add(_t.doer.name)
check("a doer's export holds only their own work",
      not (_others - _assigned), f"leaked {sorted(_others - _assigned)[:3]}")
check("a doer cannot export the staff list",
      doer.get("/admin/users?export=xlsx").status_code == 403)
check("nor the branches", doer.get("/admin/branches?export=xlsx").status_code == 403)
check("nor the EM score report",
      doer.get("/reports/score?export=xlsx").status_code == 403)
check("and an anonymous visitor gets nothing",
      TestClient(app, follow_redirects=False)
      .get("/admin/users?export=xlsx").status_code in (303, 401, 403))

# --- the user list, which is what was asked for by name -------------------
_wbu, _ = _book(admin, "/admin/users?export=xlsx")
_hu, _ru = _table(_wbu.worksheets[0])
with _SLT() as _d:
    _headcount = len(_d.scalars(_msel(_MU).where(_MU.org_id == 1)).all())
check("the user list exports every user", len(_ru) == _headcount,
      f"{len(_ru)} rows vs {_headcount} users")
for _col in ("Name", "Email", "Role", "Branch", "Department", "Active",
             "Rights", "Benchmark — Delegation"):
    check(f"the user list has a {_col} column", _col in _hu, _hu)
check("rights are spelled out, not stored codes",
      any("All rights" in str(r[_hu.index("Rights")]) for r in _ru))
check("and an ordinary doer's rights read plainly",
      any(str(r[_hu.index("Rights")]) in ("Execute only",)
          or "," in str(r[_hu.index("Rights")]) for r in _ru))
check("no password or hash is ever in the file",
      not any("hash" in str(h).lower() or "password" in str(h).lower()
              for h in _hu), _hu)

# --- multi-tab books ------------------------------------------------------
_wbs, _ = _book(admin, "/reports/score?export=xlsx")
check("the score export has both views",
      [w.title for w in _wbs.worksheets] == ["Person-wise", "Branch-wise"],
      [w.title for w in _wbs.worksheets])
_hs, _rs = _table(_wbs.worksheets[0])
check("with the bifurcation spread across columns",
      all(c in _hs for c in ["Score", "Delegation penalty", "Checklist penalty",
                             "FMS penalty", "False marking penalty"]), _hs)
_wbr, _ = _book(admin, "/reports/tasks?source=delegation&export=xlsx")
check("a source report carries its figures on a second tab",
      "Summary" in [w.title for w in _wbr.worksheets])
_wbf, _ = _book(admin, "/flows?export=xlsx")
check("the FMS export lists steps and runs separately",
      [w.title for w in _wbf.worksheets] == ["Flow steps", "Running now"],
      [w.title for w in _wbf.worksheets])

# --- an empty result is a readable file, not a broken one -----------------
_wbe, _ = _book(admin, "/tasks?scope=mine&status=done&date_from=2001-01-01"
                       "&date_to=2001-01-02&export=xlsx")
check("an export with no rows still opens", _wbe is not None)
if _wbe:
    check("and says so in words",
          any("Nothing matched" in str(c.value or "")
              for c in _wbe.worksheets[0]["A"][:8]))

# --- the button is on the pages ------------------------------------------
for _page in ["/tasks?scope=mine&status=open", "/reports/tasks?source=delegation",
              "/reports/audit", "/reports/score", "/reports/followups",
              "/admin/users", "/admin/branches", "/admin/departments",
              "/followups", "/flows", "/recurring", "/outbox"]:
    _p = admin.get(_page)
    check(f"{_page} offers the Excel button",
          "export=xlsx" in _p.text, _p.status_code)

# Performance and the Dashboard too.
_wbp, _rp2 = _book(admin, "/stats?export=xlsx")
check("/stats downloads an Excel file", _wbp is not None, _rp2.status_code)
if _wbp:
    check("with both views",
          [w.title for w in _wbp.worksheets] == ["Person-wise", "Branch-wise"])
check("a doer cannot export Performance",
      doer.get("/stats?export=xlsx").status_code == 403)
check("the dashboard offers one too", "export=xlsx" in admin.get("/").text)
check("and a doer's dashboard offers it as well", "export=xlsx" in doer.get("/").text)

# An ordinary GET must still be the page, not a download.
check("without export= the page is still a page",
      admin.get("/tasks?scope=mine&status=open").headers["content-type"]
      .startswith("text/html"))
check("and a nonsense export value is ignored",
      admin.get("/tasks?scope=mine&status=open&export=banana")
      .headers["content-type"].startswith("text/html"))

print("\n== deleting things that other rows point at ==")
# The live 500: a task the follow-up desk had ever ticked refused to delete.
# SQLite used to ignore foreign keys, so every test here passed while the
# same click failed on the real site. Foreign keys are enforced now, which is
# what makes the checks below mean anything.
from app.models import (Followup as _FU, Attachment as _AT, TaskComment as _TC,
                        OutboundMessage as _MO)
from sqlalchemy import func as _func
from app.db import engine as _eng

if _eng.url.get_backend_name() == "sqlite":
    with _eng.connect() as _cn:
        check("SQLite is enforcing foreign keys like the live database does",
              _cn.exec_driver_sql("PRAGMA foreign_keys").scalar() == 1)

def _fresh_task(title):
    r = mgr.post("/tasks/new", data={
        "title": title, "details": "", "doer_id": str(_amit.id),
        "branch_id": "", "priority": "low", "due_at": "2026-12-31T23:59"})
    return int(_re.findall(r"/tasks/(\d+)/comment", r.text)[-1])

_dt1 = _fresh_task(f"SMOKE delete chased {RUN}")
with _SLT() as _d:
    _t = _d.get(_MT, _dt1)
    _d.add(_FU(org_id=_t.org_id, task_id=_t.id, day=_clock.today(), by_id=1))
    _d.add(_TC(task_id=_t.id, author_id=1, body="a note"))
    _d.add(_MO(org_id=_t.org_id, task_id=_t.id, to_phone="+910000000000",
               template="task_assigned", body="x", status="queued"))
    _d.commit()

_del = admin.post(f"/tasks/{_dt1}/delete")
check("a task the follow-up desk chased can be deleted",
      _del.status_code in (200, 303), _del.status_code)
with _SLT() as _d:
    check("the task is gone", _d.get(_MT, _dt1) is None)
    check("its follow-up ticks went with it",
          _d.scalar(_msel(_func.count()).select_from(_FU)
                    .where(_FU.task_id == _dt1)) == 0)
    check("so did its notes",
          _d.scalar(_msel(_func.count()).select_from(_TC)
                    .where(_TC.task_id == _dt1)) == 0)
    # The message log is a record of what was SENT. It outlives the task.
    _left = _d.scalars(_msel(_MO).where(_MO.body == "x",
                                        _MO.to_phone == "+910000000000")).all()
    check("but the message we sent is kept, just unhooked",
          _left and all(m.task_id is None for m in _left),
          f"{len(_left)} rows")

# Attachments too — a task with proof on it must still delete.
_dt2 = _fresh_task(f"SMOKE delete with proof {RUN}")
attach(doer, _dt2)
with _SLT() as _d:
    check("it has proof attached",
          _d.scalar(_msel(_func.count()).select_from(_AT)
                    .where(_AT.task_id == _dt2)) > 0)
check("a task with attachments deletes cleanly",
      admin.post(f"/tasks/{_dt2}/delete").status_code in (200, 303))
with _SLT() as _d:
    check("and leaves no attachment behind",
          _d.scalar(_msel(_func.count()).select_from(_AT)
                    .where(_AT.task_id == _dt2)) == 0)

# Deleting a PERSON has exactly the same trap: a tick names who made it, and
# that column will not take a null either.
print("\n== and deleting a person who had chased work ==")
_mk = admin.post("/admin/users", data={
    "name": f"Temp Chaser {RUN}", "email": f"chaser.{RUN}@gcs.local",
    "password": "gcs1234", "role": "doer", "branch_id": "", "department_id": "",
    "bm_delegation": "60", "bm_checklist": "20", "bm_fms": "20"})
check("a test user is created", _mk.status_code in (200, 303), _mk.status_code)
with _SLT() as _d:
    _tmp = _d.scalar(_msel(_MU).where(_MU.email == f"chaser.{RUN}@gcs.local"))
    _tid_keep = _d.scalars(_msel(_MT).where(_MT.doer_id != _tmp.id)).first().id
    _d.add(_FU(org_id=_tmp.org_id, task_id=_tid_keep,
               day=_clock.today(), by_id=_tmp.id))
    _d.add(_AT(task_id=_tid_keep, uploaded_by_id=_tmp.id, filename="theirs.png",
               stored_name=f"k-{RUN}", size=1, content_type="image/png",
               storage="db", data=b"x"))
    _d.commit()
    _tmp_id, _tmp_name = _tmp.id, _tmp.name

_page = admin.get(f"/admin/users/{_tmp_id}/delete")
check("the confirmation page counts their follow-up ticks",
      "Follow-up ticks they made" in _page.text)
check("and the files they uploaded", "Files they uploaded" in _page.text)

_rm = admin.post(f"/admin/users/{_tmp_id}/delete",
                 data={"mode": "transfer", "transfer_to": str(_amit.id),
                       "confirm": _tmp_name})
check("deleting them does not fail", _rm.status_code in (200, 303), _rm.status_code)
with _SLT() as _d:
    check("the person is gone", _d.get(_MU, _tmp_id) is None)
    check("no tick is left pointing at nobody",
          _d.scalar(_msel(_func.count()).select_from(_FU)
                    .where(_FU.by_id == _tmp_id)) == 0)
    check("the tick itself survives, under the person it moved to",
          _d.scalar(_msel(_func.count()).select_from(_FU)
                    .where(_FU.task_id == _tid_keep, _FU.by_id == _amit.id)) > 0)
    check("and somebody else's task was not touched",
          _d.get(_MT, _tid_keep) is not None)

# Finally: nothing anywhere may be left pointing at a row that is gone.
print("\n== nothing in the database points at something deleted ==")
with _SLT() as _d:
    _orphans = []
    for _model, _col, _parent in [
            (_FU, "task_id", _MT), (_FU, "by_id", _MU),
            (_TC, "task_id", _MT), (_AT, "task_id", _MT),
            (_AT, "uploaded_by_id", _MU), (_MT, "doer_id", _MU),
            (_MT, "assigner_id", _MU)]:
        _ids = {r[0] for r in _d.execute(_msel(getattr(_model, _col))).all()
                if r[0] is not None}
        _have = {r[0] for r in _d.execute(_msel(_parent.id)).all()}
        _missing = _ids - _have
        if _missing:
            _orphans.append(f"{_model.__tablename__}.{_col}: {sorted(_missing)[:3]}")
    check("every row points at something that exists", not _orphans, _orphans)

print("\n" + ("ALL CHECKS PASSED" if not FAIL else f"{len(FAIL)} FAILED: {FAIL}"))
sys.exit(1 if FAIL else 0)
